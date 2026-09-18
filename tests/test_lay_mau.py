# -*- coding: utf-8 -*-
"""Cat mau N dong de chay thu, LAP LAI DUOC.

Khong co cong cu nay thi moi lan chay thu lai mot tap dong khac -> so do giua hai
lan khong so duoc voi nhau, va "hoi quy" tro thanh doan mo.
"""
import hashlib
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

GOC = Path(__file__).resolve().parents[1]
CONG_CU = GOC / "tools" / "lay_mau.py"


def _file_nguon(thu_muc: Path, so_dong: int = 50) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["part_id", "part_number", "part_description"])
    for i in range(so_dong):
        ws.append([i, "MA%03d" % i, "San pham %d" % i])
    duong_dan = thu_muc / "nguon.xlsx"
    wb.save(duong_dan)
    return duong_dan


def _chay(nguon: Path, ra: Path, rows: int, seed: int = 0):
    return subprocess.run(
        [sys.executable, str(CONG_CU), "--input", str(nguon), "--rows", str(rows),
         "--seed", str(seed), "--out", str(ra)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)


def _van_tay(duong_dan: Path) -> str:
    """Hash noi dung cac dong - de so hai lan chay co ra dung cung tap dong khong."""
    ws = openpyxl.load_workbook(duong_dan, read_only=True).worksheets[0]
    text = "\n".join("|".join(str(v) for v in hang) for hang in ws.iter_rows(values_only=True))
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def test_dung_so_dong_va_giu_nguyen_header(tmp_path):
    nguon = _file_nguon(tmp_path)
    ra = tmp_path / "mau.xlsx"
    assert _chay(nguon, ra, 10).returncode == 0
    hang = list(openpyxl.load_workbook(ra, read_only=True).worksheets[0]
                .iter_rows(values_only=True))
    assert hang[0] == ("part_id", "part_number", "part_description")
    assert len(hang) == 11                      # 1 header + 10 dong du lieu


def test_cung_seed_ra_cung_tap_dong(tmp_path):
    """Cot loi: khong lap lai duoc thi khong so duoc so do giua hai lan chay."""
    nguon = _file_nguon(tmp_path)
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _chay(nguon, a, 10, seed=7)
    _chay(nguon, b, 10, seed=7)
    assert _van_tay(a) == _van_tay(b)


def test_seed_khac_ra_tap_khac(tmp_path):
    nguon = _file_nguon(tmp_path)
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _chay(nguon, a, 10, seed=0)
    _chay(nguon, b, 10, seed=1)
    assert _van_tay(a) != _van_tay(b)


def test_xin_nhieu_hon_so_dong_co_that_thi_lay_het_khong_chet(tmp_path):
    nguon = _file_nguon(tmp_path, so_dong=5)
    ra = tmp_path / "mau.xlsx"
    ket_qua = _chay(nguon, ra, 100)
    assert ket_qua.returncode == 0
    hang = list(openpyxl.load_workbook(ra, read_only=True).worksheets[0]
                .iter_rows(values_only=True))
    assert len(hang) == 6


def test_khong_thay_file_thi_thoat_khac_0(tmp_path):
    ket_qua = _chay(tmp_path / "khong-co.xlsx", tmp_path / "ra.xlsx", 10)
    assert ket_qua.returncode != 0


def test_giu_nguyen_thu_tu_goc(tmp_path):
    """Dao thu tu dong lam kho doi chieu bang mat voi file goc."""
    nguon = _file_nguon(tmp_path)
    ra = tmp_path / "mau.xlsx"
    _chay(nguon, ra, 10)
    ids = [h[0] for h in openpyxl.load_workbook(ra, read_only=True).worksheets[0]
           .iter_rows(min_row=2, values_only=True)]
    assert ids == sorted(ids)


# ── Hold-out: xoa trang o da biet dap an de do DO CHINH XAC ──────────────────
#
# File goc 5.730 dong khong co o filter nao trong, nen khong cach nao exercise
# phan dien o trong cua mecsu-filter. Xoa trang o da biet dap an giai quyet ca
# hai: co viec cho skill lam, VA co dap an de cham.

def _chay_holdout(nguon: Path, ra: Path, rows: int, cot: str, so_o: int, seed: int = 0):
    return subprocess.run(
        [sys.executable, str(CONG_CU), "--input", str(nguon), "--rows", str(rows),
         "--seed", str(seed), "--out", str(ra), "--xoa-cot", cot, "--xoa-so-o", str(so_o)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)


def _file_co_filter(thu_muc: Path, so_dong: int = 50) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["part_id", "part_number", "mapped_filters"])
    for i in range(so_dong):
        ws.append([i, "MA%03d" % i, "Chiều dài: %d mm | Vật liệu: Thép" % (i + 10)])
    duong_dan = thu_muc / "co_filter.xlsx"
    wb.save(duong_dan)
    return duong_dan


def test_xoa_dung_so_o_va_ghi_dap_an_ra_file_ben_canh(tmp_path):
    nguon = _file_co_filter(tmp_path)
    ra = tmp_path / "mau.xlsx"
    assert _chay_holdout(nguon, ra, 20, "mapped_filters", 6).returncode == 0
    ws = openpyxl.load_workbook(ra, read_only=True).worksheets[0]
    hang = list(ws.iter_rows(values_only=True))
    cot = [str(v) for v in hang[0]].index("mapped_filters")
    trong = [h for h in hang[1:] if not str(h[cot] or "").strip()]
    assert len(trong) == 6

    import json
    dap_an = json.loads((ra.parent / (ra.stem + ".dap_an.json")).read_text(encoding="utf-8"))
    assert len(dap_an["da_xoa"]) == 6
    # Dap an phai la gia tri THAT tu file goc, khong phai chuoi rong.
    assert all(str(v).strip() for v in dap_an["da_xoa"].values())


def test_holdout_lap_lai_duoc(tmp_path):
    """Cung seed phai xoa dung nhung o do - khong thi khong so duoc hai lan chay."""
    import json
    nguon = _file_co_filter(tmp_path)
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _chay_holdout(nguon, a, 20, "mapped_filters", 6, seed=3)
    _chay_holdout(nguon, b, 20, "mapped_filters", 6, seed=3)
    da = json.loads((tmp_path / "a.dap_an.json").read_text(encoding="utf-8"))["da_xoa"]
    db = json.loads((tmp_path / "b.dap_an.json").read_text(encoding="utf-8"))["da_xoa"]
    assert da == db
    assert _van_tay(a) == _van_tay(b)


def test_xoa_cot_khong_ton_tai_thi_thoat_khac_0(tmp_path):
    """Doan ten cot sai roi im lang xoa cot khac la hong am tham."""
    nguon = _file_co_filter(tmp_path)
    ket_qua = _chay_holdout(nguon, tmp_path / "mau.xlsx", 20, "khong_co_cot_nay", 3)
    assert ket_qua.returncode != 0


def test_xoa_nhieu_hon_so_dong_thi_thoat_khac_0(tmp_path):
    nguon = _file_co_filter(tmp_path)
    ket_qua = _chay_holdout(nguon, tmp_path / "mau.xlsx", 10, "mapped_filters", 99)
    assert ket_qua.returncode != 0


# ── Lay mau THEO NHOM ────────────────────────────────────────────────────────
#
# Boc ngau nhien 100 dong tu 5.730 lam vun het nhom: do duoc 39/43 nhom chi con
# 1-4 dong, 66/100 dong khong vao duoc cum nao (can toi thieu 5 lang gieng).
# Skill hoc quy uoc TU LANG GIENG trong chinh file, nen mau vun lam no mu - va
# con so do duoc la do MAU, khong phai do SKILL.

def _chay_nhom(nguon: Path, ra: Path, rows: int, cot_nhom: str, seed: int = 0):
    return subprocess.run(
        [sys.executable, str(CONG_CU), "--input", str(nguon), "--rows", str(rows),
         "--seed", str(seed), "--out", str(ra), "--theo-nhom", cot_nhom],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)


def _file_co_nhom(thu_muc: Path) -> Path:
    """6 nhom, moi nhom 10 dong."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["part_id", "part_number", "leaf_category_name"])
    for g in range(6):
        for i in range(10):
            ws.append([g * 10 + i, "MA%d%02d" % (g, i), "Nhom %d" % g])
    duong_dan = thu_muc / "co_nhom.xlsx"
    wb.save(duong_dan)
    return duong_dan


def _dem_nhom(duong_dan: Path) -> dict:
    ws = openpyxl.load_workbook(duong_dan, read_only=True).worksheets[0]
    hang = list(ws.iter_rows(values_only=True))
    cot = [str(v) for v in hang[0]].index("leaf_category_name")
    dem = {}
    for h in hang[1:]:
        dem[h[cot]] = dem.get(h[cot], 0) + 1
    return dem


def test_lay_tron_nhom_khong_cat_le_dong(tmp_path):
    """Nhom nao duoc chon thi phai lay DU 10 dong cua no, khong lay 3 dong roi bo."""
    nguon = _file_co_nhom(tmp_path)
    ra = tmp_path / "mau.xlsx"
    assert _chay_nhom(nguon, ra, 30, "leaf_category_name").returncode == 0
    dem = _dem_nhom(ra)
    assert all(n == 10 for n in dem.values()), dem
    assert sum(dem.values()) == 30


def test_khong_du_so_dong_chinh_xac_thi_lay_qua_hoac_du_nhom_gan_nhat(tmp_path):
    """25 dong khong chia het cho nhom 10 - phai lay tron nhom, khong cat."""
    nguon = _file_co_nhom(tmp_path)
    ra = tmp_path / "mau.xlsx"
    assert _chay_nhom(nguon, ra, 25, "leaf_category_name").returncode == 0
    dem = _dem_nhom(ra)
    assert all(n == 10 for n in dem.values()), dem
    assert sum(dem.values()) in (20, 30)


def test_theo_nhom_lap_lai_duoc(tmp_path):
    nguon = _file_co_nhom(tmp_path)
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _chay_nhom(nguon, a, 30, "leaf_category_name", seed=5)
    _chay_nhom(nguon, b, 30, "leaf_category_name", seed=5)
    assert _van_tay(a) == _van_tay(b)


def test_cot_nhom_khong_ton_tai_thi_thoat_khac_0(tmp_path):
    nguon = _file_co_nhom(tmp_path)
    assert _chay_nhom(nguon, tmp_path / "mau.xlsx", 30, "khong_co").returncode != 0


# ── Gan SAI co chu dich: do xem skill co BAT dung nhung dong do ──────────────
#
# mecsu-category SOAT danh muc da gan, khong dien o trong. Nen hold-out o day la
# doi danh muc cua N dong sang mot danh muc khac CO THAT trong file, roi do xem
# skill bat dung N dong do khong (va khong bao dong gia o cac dong con lai).

def _chay_doi(nguon: Path, ra: Path, rows: int, cot: str, so_o: int, seed: int = 0):
    return subprocess.run(
        [sys.executable, str(CONG_CU), "--input", str(nguon), "--rows", str(rows),
         "--seed", str(seed), "--out", str(ra), "--doi-cot", cot, "--doi-so-o", str(so_o)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)


def _file_co_cate(thu_muc: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["part_id", "part_number", "leaf_category_name"])
    for i in range(30):
        ws.append([i, "MA%03d" % i, "Danh muc %d" % (i % 5)])
    duong_dan = thu_muc / "co_cate.xlsx"
    wb.save(duong_dan)
    return duong_dan


def _doc(duong_dan: Path, cot: str):
    ws = openpyxl.load_workbook(duong_dan, read_only=True).worksheets[0]
    hang = list(ws.iter_rows(values_only=True))
    i_c = [str(v) for v in hang[0]].index(cot)
    i_ma = [str(v) for v in hang[0]].index("part_number")
    return {str(h[i_ma]): h[i_c] for h in hang[1:]}


def test_doi_dung_so_o_va_gia_tri_moi_phai_KHAC_gia_tri_cu(tmp_path):
    import json
    nguon = _file_co_cate(tmp_path)
    ra = tmp_path / "mau.xlsx"
    assert _chay_doi(nguon, ra, 30, "leaf_category_name", 8).returncode == 0
    goc = _doc(nguon, "leaf_category_name")
    moi = _doc(ra, "leaf_category_name")
    doi = [ma for ma in moi if moi[ma] != goc[ma]]
    assert len(doi) == 8, doi
    dap_an = json.loads((tmp_path / "mau.dap_an.json").read_text(encoding="utf-8"))
    assert len(dap_an["da_doi"]) == 8
    for ma, muc in dap_an["da_doi"].items():
        assert muc["goc"] != muc["gan_sai"]
        assert moi[ma] == muc["gan_sai"]


def test_gia_tri_gan_sai_phai_CO_THAT_trong_cot_do(tmp_path):
    """Gan mot danh muc khong ton tai thi qua de bat - khong do duoc gi."""
    import json
    nguon = _file_co_cate(tmp_path)
    ra = tmp_path / "mau.xlsx"
    _chay_doi(nguon, ra, 30, "leaf_category_name", 8)
    co_that = set(_doc(nguon, "leaf_category_name").values())
    dap_an = json.loads((tmp_path / "mau.dap_an.json").read_text(encoding="utf-8"))
    for muc in dap_an["da_doi"].values():
        assert muc["gan_sai"] in co_that


def test_doi_lap_lai_duoc(tmp_path):
    nguon = _file_co_cate(tmp_path)
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _chay_doi(nguon, a, 30, "leaf_category_name", 8, seed=9)
    _chay_doi(nguon, b, 30, "leaf_category_name", 8, seed=9)
    assert _van_tay(a) == _van_tay(b)


def test_doi_cot_khong_ton_tai_thi_thoat_khac_0(tmp_path):
    nguon = _file_co_cate(tmp_path)
    assert _chay_doi(nguon, tmp_path / "m.xlsx", 30, "khong_co", 3).returncode != 0


def test_doi_nhieu_hon_so_dong_thi_thoat_khac_0(tmp_path):
    nguon = _file_co_cate(tmp_path)
    assert _chay_doi(nguon, tmp_path / "m.xlsx", 10, "leaf_category_name", 99).returncode != 0
