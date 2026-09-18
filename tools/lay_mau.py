# -*- coding: utf-8 -*-
"""Cat N dong tu mot file Excel de chay thu - LAP LAI DUOC theo seed.

    python tools/lay_mau.py --input samples/Handtools-check-cate-filter.xlsx \\
        --rows 100 --seed 0 --out jobs/mau-100.xlsx

Cung seed thi ra dung cung tap dong. Khong the thi hai lan chay thu lay hai tap
khac nhau, va moi so do so sanh giua chung deu vo nghia.

Giu nguyen HEADER va giu nguyen THU TU goc cua cac dong duoc chon: file mau phai
doi chieu duoc bang mat voi file goc.
"""
from __future__ import annotations

import argparse
import random
import sys
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--rows', type=int, default=100)
    parser.add_argument('--seed', type=int, default=0)
    parser.add_argument('--sheet', default='', help='Ten sheet. Mac dinh sheet dau tien.')
    parser.add_argument('--xoa-cot', default='',
                        help='Ten cot de XOA TRANG mot so o (hold-out). Dap an ghi ra file ben canh.')
    parser.add_argument('--xoa-so-o', type=int, default=0, help='Xoa trang bao nhieu o.')
    parser.add_argument('--doi-cot', default='',
                        help='Ten cot de GAN SAI CO CHU DICH mot so o: doi sang mot gia tri '
                             'khac CO THAT trong cung cot. Dap an ghi ra file ben canh.')
    parser.add_argument('--doi-so-o', type=int, default=0, help='Gan sai bao nhieu o.')
    parser.add_argument('--theo-nhom', default='',
                        help='Ten cot nhom. Lay TRON tung nhom cho toi khi du --rows, '
                             'thay vi boc le dong. Giu lang gieng de skill hoc duoc quy uoc.')
    args = parser.parse_args()

    if not args.input.exists():
        raise SystemExit('khong thay file: %s' % args.input)
    if args.rows < 1:
        raise SystemExit('--rows phai >= 1')

    nguon = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    ws = nguon[args.sheet] if args.sheet else nguon.worksheets[0]

    hang = list(ws.iter_rows(values_only=True))
    if not hang:
        raise SystemExit('%s khong co dong nao' % args.input)
    header, du_lieu = hang[0], hang[1:]
    if not du_lieu:
        raise SystemExit('%s chi co header, khong co dong du lieu' % args.input)

    if args.theo_nhom:
        # Boc le dong lam vun nhom: do duoc 39/43 nhom chi con 1-4 dong va 66/100 dong
        # khong vao duoc cum nao (cum can toi thieu 5 lang gieng). Skill hoc quy uoc TU
        # LANG GIENG trong chinh file, nen mau vun do chinh cai mau chu khong do skill.
        ten_cot = [str(v) if v is not None else '' for v in header]
        if args.theo_nhom not in ten_cot:
            raise SystemExit('khong co cot nhom %r. Cot that: %s' % (args.theo_nhom, ten_cot))
        cot_nhom = ten_cot.index(args.theo_nhom)
        theo_nhom: dict = {}
        for j, hang_du in enumerate(du_lieu):
            theo_nhom.setdefault(str(hang_du[cot_nhom]), []).append(j)
        ten_nhom = sorted(theo_nhom)
        random.Random(args.seed).shuffle(ten_nhom)
        chon, da_lay = [], []
        for ten in ten_nhom:
            if len(chon) >= args.rows:
                break
            chon.extend(theo_nhom[ten])             # lay TRON nhom, khong cat le
            da_lay.append(ten)
        chon.sort()
        print('nhom da lay %d/%d nhom: %s%s'
              % (len(da_lay), len(ten_nhom), ', '.join(da_lay[:6]),
                 ' ...' if len(da_lay) > 6 else ''))
    elif args.rows >= len(du_lieu):
        chon = list(range(len(du_lieu)))            # xin nhieu hon co that -> lay het
    else:
        chon = sorted(random.Random(args.seed).sample(range(len(du_lieu)), args.rows))

    lay = [list(du_lieu[i]) for i in chon]

    # Hold-out: xoa trang o DA BIET dap an. Khong phai bia du lieu - la BO du lieu,
    # va giu lai dap an de cham do chinh xac cua thu skill dien vao.
    dap_an: dict = {}
    if args.xoa_cot:
        ten_cot = [str(v) if v is not None else '' for v in header]
        if args.xoa_cot not in ten_cot:
            raise SystemExit('khong co cot %r. Cot that: %s' % (args.xoa_cot, ten_cot))
        cot = ten_cot.index(args.xoa_cot)
        if args.xoa_so_o < 1:
            raise SystemExit('--xoa-cot phai di kem --xoa-so-o >= 1')
        co_gia_tri = [j for j, h in enumerate(lay) if str(h[cot] or '').strip()]
        if args.xoa_so_o > len(co_gia_tri):
            raise SystemExit('xin xoa %d o nhung chi co %d o co gia tri trong %d dong da chon'
                             % (args.xoa_so_o, len(co_gia_tri), len(lay)))
        # Khoa nhan dien dong: uu tien part_number, khong co thi cot dau.
        cot_khoa = ten_cot.index('part_number') if 'part_number' in ten_cot else 0
        for j in sorted(random.Random(args.seed + 1).sample(co_gia_tri, args.xoa_so_o)):
            dap_an[str(lay[j][cot_khoa])] = lay[j][cot]
            lay[j][cot] = None

    # Gan SAI co chu dich: doi sang mot gia tri KHAC nhung CO THAT trong cung cot.
    # Gan mot gia tri khong ton tai thi qua de bat, khong do duoc nang luc that.
    da_doi: dict = {}
    if args.doi_cot:
        ten_cot = [str(v) if v is not None else '' for v in header]
        if args.doi_cot not in ten_cot:
            raise SystemExit('khong co cot %r. Cot that: %s' % (args.doi_cot, ten_cot))
        cot = ten_cot.index(args.doi_cot)
        if args.doi_so_o < 1:
            raise SystemExit('--doi-cot phai di kem --doi-so-o >= 1')
        co_gia_tri = [j for j, h in enumerate(lay) if str(h[cot] or '').strip()]
        if args.doi_so_o > len(co_gia_tri):
            raise SystemExit('xin doi %d o nhung chi co %d o co gia tri'
                             % (args.doi_so_o, len(co_gia_tri)))
        gia_tri_co_that = sorted({str(h[cot]).strip() for h in lay if str(h[cot] or '').strip()})
        if len(gia_tri_co_that) < 2:
            raise SystemExit('cot %r chi co mot gia tri - khong the gan sai'
                             % args.doi_cot)
        cot_khoa = ten_cot.index('part_number') if 'part_number' in ten_cot else 0
        may = random.Random(args.seed + 2)
        for j in sorted(may.sample(co_gia_tri, args.doi_so_o)):
            cu = str(lay[j][cot]).strip()
            khac = [v for v in gia_tri_co_that if v != cu]
            moi = may.choice(khac)
            da_doi[str(lay[j][cot_khoa])] = {'goc': cu, 'gan_sai': moi}
            lay[j][cot] = moi

    ra = openpyxl.Workbook()
    dich = ra.active
    dich.title = ws.title[:31]
    dich.append(list(header))
    for hang_ra in lay:
        dich.append(hang_ra)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    ra.save(args.out)

    if dap_an or da_doi:
        import json
        ben_canh = args.out.parent / (args.out.stem + '.dap_an.json')
        noi_dung = {'nguon': str(args.input), 'seed': args.seed,
                    'khoa_la_cot': ten_cot[cot_khoa]}
        if dap_an:
            noi_dung['cot'] = args.xoa_cot
            noi_dung['da_xoa'] = dap_an
        if da_doi:
            noi_dung['cot_doi'] = args.doi_cot
            noi_dung['da_doi'] = da_doi
        ben_canh.write_text(json.dumps(noi_dung, ensure_ascii=False, indent=2), encoding='utf-8')
        if dap_an:
            print('da xoa    %d o cot %r' % (len(dap_an), args.xoa_cot))
        if da_doi:
            print('da gan sai %d o cot %r' % (len(da_doi), args.doi_cot))
        print('dap an    %s' % ben_canh)

    print('nguon      %s  (%d dong du lieu)' % (args.input, len(du_lieu)))
    print('seed       %d' % args.seed)
    print('da ghi     %s  (%d dong)' % (args.out, len(chon)))
    print('dong dau   %s' % (du_lieu[chon[0]][:3],))
    print('dong cuoi  %s' % (du_lieu[chon[-1]][:3],))


if __name__ == '__main__':
    main()
