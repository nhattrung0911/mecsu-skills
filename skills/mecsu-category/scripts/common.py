"""Shared helpers for the On_web_category audit pipeline."""

from __future__ import annotations

import hashlib
import json
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import openpyxl

# Windows consoles default to cp1252 and cannot print Vietnamese category names.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

# The pipeline needs these roles filled; the workbook may name and order its columns any way it
# likes, and may carry extra columns the pipeline never touches but must preserve.
REQUIRED_ROLES = ("description", "leaf_name", "level1_name")
OPTIONAL_ROLES = ("part_id", "part_number", "sku", "leaf_id", "leaf_depth", "level1_id")
ROLES = REQUIRED_ROLES + OPTIONAL_ROLES

# Fallback for the workbook this pipeline was written against, so it runs with no schema file.
DEFAULT_SCHEMA = {
    "sheet": "On_web_category",
    "headers": [
        "part_id", "part_number", "part_description", "part_sku_code",
        "leaf_category_id", "leaf_category_name", "leaf_depth",
        "level1_category_id", "level1_category_name",
    ],
    "columns": {
        "part_id": 0, "part_number": 1, "description": 2, "sku": 3,
        "leaf_id": 4, "leaf_name": 5, "leaf_depth": 6,
        "level1_id": 7, "level1_name": 8,
    },
    "detected_by": "built-in default",
}


@dataclass(frozen=True)
class Row:
    """One workbook row. Named fields are the roles the pipeline reasons about; `raw` is every
    original cell, so writing a workbook back out never drops a column it did not understand."""

    part_id: object
    part_number: object
    description: str
    sku: object
    leaf_id: object
    leaf_name: str
    leaf_depth: object
    level1_id: object
    level1_name: str
    raw: tuple = ()

    def as_tuple(self) -> tuple:
        return self.raw

    def replacing(self, **roles: object) -> tuple:
        """The original row with the named roles swapped, every other column untouched."""
        cells = list(self.raw)
        for role, value in roles.items():
            index = COL.get(role)
            if index is not None and index < len(cells):
                cells[index] = value
        return tuple(cells)


def job_root() -> Path:
    """Directory that holds `jobs/`.

    These scripts also ship inside a plugin, where the enclosing repo is gone. Resolution order:
    the ONCHECK_JOBS environment variable, then the nearest ancestor containing `jobs/`, then the
    current directory. Every default path is derived from this, so nothing is hardcoded to one repo.
    """
    override = os.environ.get("ONCHECK_JOBS")
    if override:
        return Path(override)
    for parent in Path(__file__).resolve().parents:
        if (parent / "jobs").is_dir():
            return parent
    return Path.cwd()


def schema_path() -> Path:
    override = os.environ.get("ONCHECK_SCHEMA")
    return Path(override) if override else job_root() / "jobs" / "oncheck" / "schema.json"


def load_schema() -> dict:
    """The active column mapping. Absent schema.json means the workbook this was written against."""
    path = schema_path()
    if not path.exists():
        return DEFAULT_SCHEMA
    schema = json.loads(path.read_text(encoding="utf-8"))
    missing = [role for role in REQUIRED_ROLES if role not in schema.get("columns", {})]
    if missing:
        raise SystemExit(f"{path.name} does not map required role(s): {missing}")
    return schema


SCHEMA = load_schema()
SHEET = SCHEMA["sheet"]
HEADERS = list(SCHEMA["headers"])
COL = dict(SCHEMA["columns"])


def fold(value: object) -> str:
    """Lowercase, strip Vietnamese diacritics, collapse to alphanumeric words."""
    text = str(value if value is not None else "").lower().replace("đ", "d")
    text = unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def verdict_key(description: object, leaf_name: object) -> str:
    """Stable key for one (description pattern, leaf) pair — what a verdict belongs to.

    Verdicts used to be filed under `pair_id`, which rule_check.py hands out in the order
    pairs are met. Any change to the input renumbers them, so a resumed run re-attached old
    answers to different products. The pair's own content is the only key that survives.
    """
    payload = json.dumps([pattern(description), fold(leaf_name)], ensure_ascii=False)
    return hashlib.sha1(payload.encode("utf-8")).hexdigest()[:16]


def pair_key(pair: dict) -> str:
    """The key a pair's verdict is filed under. Old queues without `key_hash` still resolve."""
    return pair.get("key_hash") or verdict_key(pair["sample_description"], pair["leaf_category_name"])


def pattern(description: object) -> str:
    """Description with size/code tokens removed, so variants of one product collapse."""
    words = [word for word in fold(description).split() if not any(c.isdigit() for c in word)]
    return " ".join(words)


def read_rows(path: Path, sheet: str = SHEET) -> list[Row]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet]
    width = len(HEADERS)
    identity = COL.get("part_id", COL["description"])
    # Not every workbook has an id column. Row order is never changed, so the ordinal is a sound
    # stand-in for verifying that a written row is the row it came from.
    has_id = "part_id" in COL
    position = 0
    rows: list[Row] = []
    for index, values in enumerate(worksheet.iter_rows(min_row=1, values_only=True)):
        if index == 0:
            actual = [str(v).strip() if v is not None else "" for v in values[:width]]
            if actual != HEADERS:
                raise SystemExit(
                    f"Header of {path.name} does not match the active schema.\n"
                    f"  expected: {HEADERS}\n  found:    {actual}\n"
                    f"Run detect_schema.py against this workbook first."
                )
            continue
        cells = tuple(values[:width])
        if cells[identity] is None:
            continue

        def cell(role: str) -> object:
            position = COL.get(role)
            return cells[position] if position is not None and position < len(cells) else None

        position += 1
        rows.append(
            Row(
                part_id=cell("part_id") if has_id else f"row{position}",
                part_number=cell("part_number"),
                description=cell("description"),
                sku=cell("sku"),
                leaf_id=cell("leaf_id"),
                leaf_name=cell("leaf_name"),
                leaf_depth=cell("leaf_depth"),
                level1_id=cell("level1_id"),
                level1_name=cell("level1_name"),
                raw=cells,
            )
        )
    workbook.close()
    return rows


def write_rows(path: Path, rows: list[tuple], headers: list[str], sheet: str = "data") -> None:
    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet(sheet)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def safe_name(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", fold(value)).strip("-")
    return slug or "unknown"
