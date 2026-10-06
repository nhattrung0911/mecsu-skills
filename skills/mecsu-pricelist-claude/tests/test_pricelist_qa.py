import copy
import pytest
from openpyxl import load_workbook
from pl_common import load_theme
from pl_build import build
def _T():
    return load_theme({"image": {"required": False}})


from pl_qa import qa
from test_pricelist_build import make_doc, img_path  # noqa: F401


def _build(tmp_path, doc, name="o.xlsx"):
    out = tmp_path / name
    build(doc, str(out), load_theme())
    return out


def test_qa_clean_doc_passes(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    r = qa(str(_build(tmp_path, doc)), doc, _T())
    assert r["errors"] == []
    assert r["stats"]["pages"] == 2 and r["stats"]["data_rows"] == 4
    assert any("checking" in w for w in r["warnings"])  # non-numeric price


def test_qa_detects_overlap(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    doc["cards"][1]["_geo"]["col"] = 10
    r = qa(str(_build(tmp_path, doc)), doc, _T())
    assert any("overlap" in e for e in r["errors"])


def test_qa_detects_bad_page_code(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    out = _build(tmp_path, doc)
    wb = load_workbook(out)
    wb.worksheets[0]["AA154"] = "STA - 05"
    wb.save(out)
    r = qa(str(out), doc, _T())
    assert any("page code" in e for e in r["errors"])


def test_qa_flags_non_numeric_price(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    r = qa(str(_build(tmp_path, doc)), None, load_theme())
    assert r["errors"] == []
    assert any("0059504" in w and "non-numeric" in w for w in r["warnings"])


def test_qa_detects_missing_sheet2_code_and_non_formula(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    doc["prices"] = doc["prices"][:1]
    out = _build(tmp_path, doc)
    r = qa(str(out), None, load_theme())
    assert any("0059516" in e and "Sheet2" in e for e in r["errors"])
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=VLOOKUP"):
                c.value = 123
                break
        else:
            continue
        break
    wb.save(out)
    assert any("VLOOKUP" in e for e in qa(str(out), None, load_theme())["errors"])


def test_qa_warns_on_card_warnings_and_review_image(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    doc["cards"][0]["_warnings"] = ["image missing"]
    doc["cards"][0]["image"]["status"] = "review"
    r = qa(str(_build(tmp_path, doc)), doc, _T())
    assert any("image missing" in w for w in r["warnings"]) and any("review" in w for w in r["warnings"])


def test_qa_row_count_mismatch(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    out = _build(tmp_path, doc)
    doc["cards"][0]["rows"].append({"code": "x", "len": 1, "order": "0059504"})
    r = qa(str(out), doc, _T())
    assert any("row count" in e for e in r["errors"])


def test_qa_span_column_mismatch(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    out = _build(tmp_path, doc)
    doc["cards"][0]["_geo"]["spans"] = [4, 3, 10]
    assert any("spans for" in e for e in qa(str(out), doc, _T())["errors"])


# ---- uniformity metrics
from pl_qa import uniformity, compare_table  # noqa: E402


def test_uniformity_metrics(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    out = _build(tmp_path, doc)
    u = uniformity(str(out), doc, _T())
    assert u["errors"] == []
    assert set(u["page_fill"]) == {0, 1} and 0 < u["page_fill"][0] < 1
    assert u["image_px"]["n"] == 1 and u["desc_rows"] == {3: 3}
    assert u["title_len"]["max"] == 8 and u["labels"]["Mã Hãng"] == 3
    assert u["fonts"] == {"Calibri"} and u["formats"] <= {"General", "@", "#,##0"}
    assert any("fill" in w for w in u["warnings"]) and any("no image" in w for w in u["warnings"])


def test_uniformity_flags_offtheme_font_and_format(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    out = _build(tmp_path, doc)
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    ws["D7"].font = ws["D7"].font.copy(name="Arial", sz=13)
    ws["C13"].number_format = "0.000"
    wb.save(out)
    u = uniformity(str(out), doc, _T())
    assert any("Arial" in e for e in u["errors"]) and any("13" in e for e in u["errors"])
    assert any("0.000" in e for e in u["errors"])
    assert any("Arial" in e for e in qa(str(out), doc, _T())["errors"])


def test_compare_table_detects_style_difference(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    a = _build(tmp_path, doc, "a.xlsx")
    b = _build(tmp_path, doc, "b.xlsx")
    wb = load_workbook(b)
    ws = wb.worksheets[0]
    ws["D7"].font = ws["D7"].font.copy(sz=13)
    wb.save(b)
    th = load_theme()
    text, ok = compare_table([(str(a), doc), (str(a), doc)], th)
    assert ok and "Y" in text
    text, ok = compare_table([(str(a), doc), (str(b), doc)], th)
    assert not ok and "N" in text


def test_uniformity_skips_continuation_chunks(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    u = uniformity(str(_build(tmp_path, doc)), doc, _T())
    assert not any("a#2" in w for w in u["warnings"])  # cont chunk has no image by design
    assert any("card b: no image" in w for w in u["warnings"])


def test_qa_drawing_counted_and_bounds(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    c = doc["cards"][0]
    c["drawing"] = {"path": str(img_path)}
    c["_geo"]["drawing"] = {"col": 14, "row": 8, "col_off_emu": 0, "row_off_emu": 0, "cx_emu": 500000, "cy_emu": 400000}
    out = _build(tmp_path, doc)
    r = qa(str(out), doc, _T())
    assert r["errors"] == [] and r["stats"]["drawings"] == 1
    assert r["stats"]["uniformity"]["image_px"]["n"] == 1
    c["_geo"]["drawing"]["col"] = 40
    assert any("drawing anchor outside" in e for e in qa(str(out), doc, _T())["errors"])


def test_qa_scans_mixed_height_rows(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    doc["cards"] = doc["cards"][:1]
    doc["cards"][0]["_geo"]["image"] = None
    doc["cards"][0]["image"] = None
    doc["cards"][0]["rows"] = [{"code": f"84-{i}", "len": i, "order": f"00595{i:02d}"} for i in range(1, 4)]
    doc["cards"][0]["_geo"]["height"] = 2 + 3 + 1 + 2 + 2 + 3 + 1
    for c in doc["cards"]:
        c["_geo"]["page"] = 0
    doc["cards"][0]["data_row_heights"] = None
    out = _build(tmp_path, doc)
    wb = load_workbook(out)
    ws = wb.worksheets[0]
    # rebuild data area by hand: rows of height 1, 2, 3 stacked (code merged over its own height)
    base = 13
    for r in range(base, base + 6):
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row == r:
                ws.unmerge_cells(str(rng))
    ws["J11"].value = "Mã đặt hàng"  # real files vary the label's case
    heights = [1, 2, 3]
    r = base
    for h, row in zip(heights, doc["cards"][0]["rows"]):
        for col, span in zip((3, 7, 10, 15), (4, 3, 5, 5)):
            if h > 1 and col != 10:  # order column left unmerged: only the code column defines row height
                ws.merge_cells(start_row=r, start_column=col, end_row=r + h - 1, end_column=col + span - 1)
        ws.cell(r, 3).value = row["code"]
        ws.cell(r, 7).value = row["len"]
        ws.cell(r, 10).value = row["order"]
        ws.cell(r, 10).number_format = "@"
        ws.cell(r, 15).value = f"=VLOOKUP(J{r},Sheet2!$A:$B,2,0)"
        r += h
    wb.save(out)
    from pl_qa import _scan_cards
    wb = load_workbook(out)
    cards = _scan_cards(wb.worksheets[0], load_theme())
    assert [len(c["rows"]) for c in cards] == [3]


def test_missing_image_is_error_by_default(tmp_path, img_path):
    doc = make_doc(tmp_path, img_path)
    first = next(c for c in doc["cards"] if not c["_geo"].get("cont"))
    first["image"], first["_geo"]["image"] = None, None
    out = _build(tmp_path, doc)
    assert any("no image" in e for e in uniformity(str(out), doc, load_theme())["errors"])
    assert not any("no image" in e for e in uniformity(str(out), doc, _T())["errors"])
