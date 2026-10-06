import pytest
from PIL import Image

import pl_build
import pl_extract as px
import pl_layout
from pl_common import load_theme


def _doc(tmp_path):
    Image.new("RGBA", (120, 80), (200, 30, 30, 255)).save(tmp_path / "a.png")
    cols = [{"key": "code"}, {"key": "l", "label": "L\n(mm)"}, {"key": "b", "label": "B\n(mm)"},
            {"key": "order"}, {"key": "price"}]
    return {
        "meta": {"brand": "TESTB", "page_code": "TST"},
        "prices": [{"order_code": "0000001", "price": 1000}, {"order_code": "0000002", "price": 2000},
                   {"order_code": "0000003", "price": 3000}],
        "cards": [
            {"id": "kim-1", "title": "KÌM MỎ NHỌN", "lines": ["Vật Liệu: Thép", "Tay Cầm Bọc Nhựa"],
             "legend": ["L: Chiều Dài Kìm", "B: Chiều Rộng"], "note": "Ghi Chú Thử Nghiệm",
             "image": {"path": str(tmp_path / "a.png"), "status": "ok"}, "columns": cols,
             "rows": [{"code": "84-000", "l": 150, "b": 44, "order": "0000001"},
                      {"code": "84-002", "l": 200, "b": 55, "order": "0000002"}]},
            {"id": "kim-2", "title": "KÌM CẮT", "lines": [], "legend": [], "note": None, "image": None,
             "columns": [{"key": "code"}, {"key": "size", "label": "Size\n(inch)"}, {"key": "order"}, {"key": "price"}],
             "rows": [{"code": "X-1", "size": 8, "order": "0000003"}]},
        ],
    }


def test_norm_label():
    assert px.norm_label("Size (mm)") == "Size\n(mm)"
    assert px.norm_label("Chiều Dài (in)") == "Chiều Dài\n(inch)"
    assert px.norm_label("L \n(mm)") == "L\n(mm)"
    assert px.norm_label("Hàm") == "Hàm"


def test_build_extract_roundtrip(tmp_path):
    theme = load_theme()
    placed = pl_layout.layout(_doc(tmp_path), theme, base_dir=tmp_path)
    xlsx = str(tmp_path / "t.xlsx")
    pl_build.build(placed, xlsx, theme, base_dir=tmp_path)
    got = px.extract(xlsx, [1], work=str(tmp_path / "w"))
    assert got["meta"]["brand"] == "TESTB" and got["meta"]["page_code"] == "TST"
    assert len(got["cards"]) == 2
    by = {c["title"]: c for c in got["cards"]}
    for orig in placed["cards"]:
        c = by[orig["title"]]
        g = orig["_geo"]
        assert c["place"] == {"page": 0, "col": g["col"], "row": g["row"], "locked": True}
        assert c["lines"] == orig["lines"] and c["legend"] == orig["legend"] and c["note"] == orig["note"]
        assert [r["code"] for r in c["rows"]] == [r["code"] for r in orig["rows"]]
        assert [r["order"] for r in c["rows"]] == [r["order"] for r in orig["rows"]]
        assert [x.get("label") for x in c["columns"][1:-2]] == [x.get("label") for x in orig["columns"][1:-2]]
        assert bool(c["image"]) == bool(orig["image"])
    assert {p["order_code"]: p["price"] for p in got["prices"]} == {"0000001": 1000, "0000002": 2000, "0000003": 3000}
    assert (tmp_path / "w" / "img" / "kim-mo-nhon.png").exists()


def _handmade(path):
    """1 page, 1 card: 2-row-high data rows, lower-case 'Mã đặt hàng', formula order cell, price sheet 'Giá'."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["S1"] = "BẢNG GIÁ ZZZ"
    ws["AA77"] = "ZZZ - 1"
    ws["C5"] = "KÌM THỬ"
    ws["C5"].font = Font(bold=True)
    ws.merge_cells("C5:J6")
    ws["D7"] = "Chú thích nghiêng"
    ws["D7"].font = Font(italic=True)
    ws["C10"] = "Mã Hãng"
    ws.merge_cells("C10:E11")
    ws["F10"] = "L (mm)"
    ws.merge_cells("F10:G11")
    ws["H10"] = "Mã đặt hàng"
    ws.merge_cells("H10:J11")
    ws["K10"] = "Giá/Cái Chưa VAT"
    ws.merge_cells("K10:M11")
    for i, (code, order) in enumerate([("A-1", "0000011"), ("A-2", None), ("A-3", "0000013")]):
        r = 12 + 2 * i
        ws.cell(r, 3, code)
        ws.merge_cells(start_row=r, start_column=3, end_row=r + 1, end_column=5)
        ws.cell(r, 6, 100 + i)
        ws.cell(r, 8, order if order else "=VLOOKUP(1,Z:Z,1,0)")
        ws.merge_cells(start_row=r, start_column=8, end_row=r + 1, end_column=10)
    s = wb.create_sheet("Giá")
    s.append(["order_id", "varin_price"])
    s.append(["0000011", 10])
    s.append(["0000013", 30])
    wb.save(path)


def test_extract_edge_cases(tmp_path):
    p = str(tmp_path / "h.xlsx")
    _handmade(p)
    d = px.extract(p, [1], reflow=True)
    c = d["cards"][0]
    assert [r["code"] for r in c["rows"]] == ["A-1", "A-2", "A-3"]  # 2-high rows are not lost
    assert c["rows"][0]["order"] == "0000011"
    assert {x["order_code"]: x["price"] for x in d["prices"]}["0000013"] == 30  # price sheet auto-detected
    assert c["note"] == "Chú thích nghiêng"  # italic-only caption
    assert "place" not in c and c["width"] == "auto" and all("span" not in x for x in c["columns"])
    assert c["columns"][1]["label"] == "L\n(mm)"
    assert d["meta"]["page_code"] == "ZZZ"


def test_classify_photo_vs_drawing():
    line = Image.new("RGBA", (200, 100), (0, 0, 0, 0))
    for x in range(200):
        for y in range(0, 100, 5):
            line.putpixel((x, y), (0, 0, 0, 120))
    body = Image.new("RGBA", (200, 100), (0, 0, 0, 0))
    body.paste((120, 120, 130, 255), (40, 20, 160, 80))
    assert px._classify(line) == "drawing"
    assert px._classify(body) == "photo"


def test_variety_pair_and_single():
    from PIL import ImageDraw
    thin = Image.new("RGBA", (300, 60), (0, 0, 0, 0))
    ImageDraw.Draw(thin).line([(0, 30), (300, 30)], fill=(90, 90, 100, 255), width=6)  # thin wrench-like photo
    photo = Image.new("RGBA", (200, 120), (0, 0, 0, 0))
    for x in range(40, 160):
        for y in range(20, 100):
            v = 60 + (x * 2 + y) % 160
            photo.putpixel((x, y), (v, v, min(255, v + 10), 255))
    art = Image.new("RGBA", (200, 120), (0, 0, 0, 0))
    ImageDraw.Draw(art).rectangle([40, 20, 160, 100], outline=(0, 0, 0, 255), width=1)
    assert px._variety(art) < 0.7 * px._variety(photo)


def test_single_picture_is_image_pair_is_split(tmp_path):
    import openpyxl
    from openpyxl.drawing.image import Image as XImage
    from PIL import ImageDraw
    p = str(tmp_path / "h.xlsx")
    _handmade(p)
    thin = Image.new("RGBA", (300, 40), (0, 0, 0, 0))
    ImageDraw.Draw(thin).line([(0, 20), (300, 20)], fill=(90, 90, 100, 255), width=3)
    photo = Image.new("RGBA", (200, 120), (0, 0, 0, 0))
    for x in range(40, 160):
        for y in range(20, 100):
            v = 60 + (x * 2 + y) % 160
            photo.putpixel((x, y), (v, v, v, 255))
    art = Image.new("RGBA", (200, 120), (0, 0, 0, 0))
    ImageDraw.Draw(art).rectangle([40, 20, 160, 100], outline=(0, 0, 0, 255), width=1)
    for name, im in (("thin", thin), ("photo", photo), ("art", art)):
        im.save(tmp_path / f"{name}.png")

    def with_pics(names, out):
        wb = openpyxl.load_workbook(p)
        ws = wb["Sheet1"]
        for i, n in enumerate(names):
            x = XImage(str(tmp_path / f"{n}.png"))
            ws.add_image(x, f"{'K' if i else 'H'}5")
        wb.save(out)
        return px.extract(out, [1], reflow=True)["cards"][0]

    c = with_pics(["thin"], str(tmp_path / "one.xlsx"))
    assert c["image"] and not c.get("drawing")
    c = with_pics(["photo", "art"], str(tmp_path / "two.xlsx"))
    assert c["image"] and c["drawing"]
