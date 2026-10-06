"""Extract a doc.json (cards with locked places) from an existing price-list xlsx in the Varin standard.

    python pl_extract.py file.xlsx --pages 1 -o doc.json --work DIR [--brand-logo PATH] [--reflow]

--reflow: no locked places and no fixed column spans (width auto) so pl_layout re-lays everything out.
Default keeps places locked (golden test).

Pages are 1-based. Card rectangle = header columns span; title = bold text above the header inside the card
columns; lines = non-bold texts between title and the VAT row; italic-only caption cells -> note/lines;
pictures anchored inside the card are classified photo (-> image) or technical drawing (-> drawing).
"""
from __future__ import annotations

import argparse
import io
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

import pathlib as _pathlib, sys as _sys  # noqa: E401
_sys.path.insert(0, str(_pathlib.Path(__file__).resolve().parent))

from pl_common import col_letter, column_role, load_theme, norm_order_code, save_doc

_LEGEND_RE = re.compile(r"^[^\s:]{1,4}\s*:")
_UNIT_RE = re.compile(r"^(.*?)(?:[\s]+\(\s*|\()(mm²|mm2|mm|cm|inch|in|kN|rpm|N\.m|kg|gam|g|m|V|A)\s*\)\s*$|^(.*?)[\s]+(mm²|mm|cm|inch|in|kN|rpm)\s*$", re.I | re.S)
_UNIT_CANON = {"in": "inch", "inch": "inch", "mm2": "mm²", "kn": "kN", "gam": "g"}


def norm_label(label: str) -> str:
    """Standard header label: '<Name>\\n(<unit>)' with a literal newline; whitespace collapsed."""
    s = str(label).replace("\r", "").strip()
    if not s:
        return s
    m = _UNIT_RE.match(s)
    flat = re.sub(r"\s+", " ", s)
    if m and (m.group(1) or m.group(3) or "").strip():
        unit = m.group(2) or m.group(4)
        unit = _UNIT_CANON.get(unit.lower(), unit if unit in ("kN", "N.m", "rpm", "V", "A") else unit.lower())
        name = re.sub(r"\s+", " ", m.group(1) or m.group(3)).strip()
        return f"{name}\n({unit})"
    return flat if "\n" not in s else "\n".join(re.sub(r"\s+", " ", p).strip() for p in s.split("\n") if p.strip())


def _src_rects(xlsx: str) -> dict:
    """{(col, row, colOff, rowOff): (l, t, r, b)} crop fractions per picture anchor (openpyxl drops srcRect)."""
    import zipfile
    try:
        with zipfile.ZipFile(xlsx) as z:
            name = next(x for x in z.namelist() if re.search(r"xl/drawings/drawing\d+\.xml$", x))
            xml = z.read(name).decode("utf-8")
    except Exception:
        return {}
    out = {}
    for blk in re.findall(r"<xdr:(?:one|two)CellAnchor.*?</xdr:(?:one|two)CellAnchor>", xml, re.S):
        if "<xdr:pic>" not in blk:
            continue
        fr = re.search(r"<xdr:from>(.*?)</xdr:from>", blk, re.S).group(1)
        key = tuple(int(re.search(rf"<xdr:{k}>(-?\d+)</xdr:{k}>", fr).group(1)) for k in ("col", "row", "colOff", "rowOff"))
        m = re.search(r"<a:srcRect([^>]*)/?>", blk)
        at = dict(re.findall(r'(\w)="(-?\d+)"', m.group(1))) if m else {}
        out.setdefault(key, tuple(int(at.get(k, 0)) / 100000 for k in "ltrb"))
    return out


def _slug(s: str) -> str:
    s = unicodedata.normalize("NFD", str(s).replace("đ", "d").replace("Đ", "D"))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def _txt(v) -> str:
    return re.sub(r"[ \t]+", " ", str(v)).strip() if v is not None else ""


def _fold(s) -> str:
    """lower-case, no diacritics, '_' -> ' ', collapsed spaces (label comparison)."""
    s = unicodedata.normalize("NFD", str(s or "").replace("đ", "d").replace("Đ", "D"))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.lower().replace("_", " ")).strip()


_ORDER_HDR = {"ma dat hang", "order id", "orderid", "part id", "ma don hang"}


def _price_table(wb, theme, skip_sheet) -> dict:
    """{order_code: price} from the first sheet whose header row has an order-code column and a price column."""
    names = [n for n in wb.sheetnames if n != skip_sheet]
    names.sort(key=lambda n: n != theme["sheet2"]["name"])
    for n in names:
        ws = wb[n]
        for hr in range(1, 4):
            heads = [_fold(c.value) for c in ws[hr]] if ws.max_row >= hr else []
            oc = next((i for i, h in enumerate(heads) if h in _ORDER_HDR), None)
            pc = next((i for i, h in enumerate(heads) if "price" in h or h.startswith("gia")), None)
            if oc is None or pc is None:
                continue
            out = {}
            for row in ws.iter_rows(min_row=hr + 1, values_only=True):
                code = norm_order_code(row[oc]) if oc < len(row) else None
                if code and code not in out:
                    out[code] = row[pc] if pc < len(row) else None
            if out:
                return out
    return {}


def _classify(img) -> str:
    """'drawing' (technical line art) or 'photo'. img: PIL RGBA.

    Cut-out PNGs: line art is anti-aliased thin lines, so almost no pixel is fully opaque (< 3 %) while
    a photo has a solid body. Opaque-background images: low colour saturation and >= 70 % near-white."""
    t = img.copy()
    t.thumbnail((200, 200))
    px = list(t.getdata())
    n = len(px) or 1
    transp = sum(1 for q in px if q[3] < 10) / n
    if transp > 0.2:
        opaque = sum(1 for q in px if q[3] > 245) / n
        semi = sum(1 for q in px if 10 <= q[3] <= 245) / n
        return "drawing" if opaque < 0.03 and semi >= 0.08 else "photo"
    white = sum(1 for r, g, b, a in px if min(r, g, b) > 235)
    objs = [q[:3] for q in px if min(q[:3]) <= 235]
    if not objs:
        return "photo"
    sat = sum((max(q) - min(q)) / max(q) if max(q) else 0 for q in objs) / len(objs)
    return "drawing" if sat < 0.12 and white / n >= 0.70 else "photo"


def _variety(img) -> int:
    """Tonal variety: occupied luminance levels (of 32) among visible pixels, weighted by >= 0.5 % share."""
    t = img.copy()
    t.thumbnail((200, 200))
    lum = [(q[0] + q[1] + q[2]) // 24 for q in t.getdata() if q[3] >= 10 and min(q[:3]) <= 235]
    if not lum:
        return 0
    cnt = {}
    for v in lum:
        cnt[v] = cnt.get(v, 0) + 1
    return sum(1 for n in cnt.values() if n / len(lum) >= 0.005)


def _val(v):
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _nearest_width(w: int, theme: dict) -> int:
    return min(theme["lanes"]["widths"], key=lambda c: (abs(c - w), c))


def _page_info(ws, off: int, theme: dict) -> tuple[str | None, str | None]:
    """(brand, page_code letters) from the page frame."""
    brand = code = None
    t = ws[theme["header"]["title"]["range"].split(":")[0]]
    m = re.match(r"\s*BẢNG GIÁ\s+(.+)", _txt(ws.cell(t.row + off, t.column).value), re.I)
    if m:
        brand = m.group(1).strip()
    pc = ws[theme["footer"]["page_code"]["range"].split(":")[0]]
    m = re.match(r"\s*([A-Za-z]{2,4})\b", _txt(ws.cell(pc.row + off, pc.column).value))
    if m:
        code = m.group(1).upper()
    return brand, code


def _headers(ws, off: int, theme: dict):
    P = theme["page"]
    label = _fold(theme["table"]["labels"]["code"])
    out = []
    for r in range(off + P["body_top"], off + P["body_bottom"] + 1):
        for c in range(P["body_left"], P["body_right"] + 1):
            if _fold(ws.cell(r, c).value) == label:
                out.append((r, c))
    return out


def _extract_card(ws, merges, bold_ital, hr, hc, off, theme, page_idx, prices_src, used_ids):
    P, C, T = theme["page"], theme["card"], theme["table"]
    # header labels
    hh = merges[(hr, hc)][3] - hr + 1 if (hr, hc) in merges else 1
    cols_raw, c = [], hc
    while c <= P["cols"]:
        v = ws.cell(hr, c).value
        if v is None:
            break
        m = merges.get((hr, c))
        w = (m[2] - c + 1) if m else 1
        cols_raw.append((str(v), c, w))
        c += w
    c_end = c - 1
    # columns
    columns, seen = [], {}
    for i, (lab, cs, w) in enumerate(cols_raw):
        flat = _txt(lab)
        if i == 0:
            columns.append({"key": "code", "role": "code", "span": w})
        elif _fold(flat) == _fold(T["labels"]["order_code"]):
            columns.append({"key": "order", "role": "order_code", "span": w})
        elif i == len(cols_raw) - 1 and _fold(flat).startswith("gia"):
            columns.append({"key": "price", "role": "price", "span": w})
        else:
            k = _slug(flat) or f"c{i}"
            seen[k] = seen.get(k, 0) + 1
            if seen[k] > 1:
                k = f"{k}-{seen[k]}"
            columns.append({"key": k, "label": norm_label(lab), "span": w})
    # data rows
    cm = merges.get((hr + hh, hc))
    drh = (cm[3] - cm[0] + 1) if cm else 1
    rows, r = [], hr + hh
    start_of = {cs: col for (_, cs, _), col in zip(cols_raw, columns)}
    while r <= off + P["body_bottom"]:
        code = ws.cell(r, hc).value
        if code is None or _txt(code) == "" or bold_ital(ws.cell(r, hc)):
            break
        row = {}
        for (_, cs, _), col in zip(cols_raw, columns):
            role = column_role(col)
            if role == "price":
                continue
            v = ws.cell(r, cs).value
            if role == "order_code":
                row[col["key"]] = norm_order_code(v) or _txt(v)
            elif role == "code":
                row[col["key"]] = _txt(v)
            else:
                row[col["key"]] = "" if v is None else _val(v)
        rows.append(row)
        mm = merges.get((r, hc))
        r += (mm[3] - mm[0] + 1) if mm else 1
    last_row = r - 1
    # title / lines above the header (row hr-1 is the VAT row)
    title, title_top, lines, captions = None, None, [], []
    rr = hr - 2
    while rr >= max(off + P["body_top"], hr - 14) and title is None:
        for cc in range(hc, c_end + 1):
            v = ws.cell(rr, cc).value
            t = _txt(v)
            if not t or t.startswith("VAT"):
                continue
            f = ws.cell(rr, cc).font
            if f.b and not f.i:
                title, title_top = t, rr
            elif f.i and not f.b:
                captions.append(t)
            else:
                lines.append(t)
            break
        rr -= 1
    captions.reverse()
    lines.reverse()
    if title is None:
        title, title_top = "KHÔNG TIÊU ĐỀ", hr - 1
    # note under the table
    note = None
    nr = r
    if nr <= off + P["body_bottom"]:
        for cc in range(hc, c_end + 1):
            cell = ws.cell(nr, cc)
            if cell.value and bold_ital(cell):
                note = _txt(cell.value)
                break
    for rr2 in range(nr, min(nr + 3, off + P["body_bottom"] + 1)):  # italic-only captions below the table
        for cc in range(hc, c_end + 1):
            cell = ws.cell(rr2, cc)
            if cell.value and cell.font.i and not cell.font.b and _txt(cell.value):
                captions.append(_txt(cell.value))
                break
    long_caps = []
    for t in captions:
        if not note:
            note = t
        elif len(t) <= 28:
            lines.append(t)
        else:
            long_caps.append(t)
    if long_caps:
        note = " ".join([note or ""] + long_caps).strip()
    legend = [x for x in lines if _LEGEND_RE.match(x)]
    lines = [x for x in lines if not _LEGEND_RE.match(x)]
    width = _nearest_width(c_end - hc + 1, theme)
    cid = _slug(title) or "card"
    base, n = cid, 1
    while cid in used_ids:
        n += 1
        cid = f"{base}-{n}"
    used_ids.add(cid)
    card = {"id": cid, "group": title, "title": title.upper(), "title_source": "data", "lines": lines, "legend": legend, "note": note,
            "image": None, "columns": columns, "rows": rows, "width": width,
            "place": {"page": page_idx, "col": hc, "row": title_top - off, "locked": True}}
    if drh >= 3:
        card["data_row_height"] = 3
    card["_rect"] = (hc, c_end, title_top - off, last_row - off)  # temporary, removed by extract()
    return card


def extract(xlsx: str, pages: list[int], work: str | None = None, brand_logo: str | None = None,
            keep_rect: bool = False, reflow: bool = False) -> dict:
    theme = load_theme()
    P = theme["page"]
    wb = load_workbook(xlsx, data_only=True)  # cached values: formula order codes / external links resolve
    ws = wb[theme["sheet1_name"]] if theme["sheet1_name"] in wb.sheetnames else wb.worksheets[0]
    merges = {(r.min_row, r.min_col): (r.min_row, r.min_col, r.max_col, r.max_row) for r in ws.merged_cells.ranges}

    def bold_ital(cell):
        return bool(cell.font.b and cell.font.i)

    prices = _price_table(wb, theme, ws.title)
    imgs = list(ws._images)
    crops = _src_rects(xlsx)

    def crop_of(im):
        f = im.anchor._from
        return crops.get((f.col, f.row, f.colOff, f.rowOff), (0, 0, 0, 0))

    img_dir = Path(work) / "img" if work else None
    if img_dir:
        img_dir.mkdir(parents=True, exist_ok=True)
    brand = page_code = None
    cards, used = [], set()
    for page_idx, pg in enumerate(pages):
        off = (pg - 1) * P["rows"]
        b, pc = _page_info(ws, off, theme)
        brand, page_code = brand or b, page_code or pc
        for hr, hc in _headers(ws, off, theme):
            cards.append(_extract_card(ws, merges, bold_ital, hr, hc, off, theme, page_idx, prices, used))
            cards[-1]["_off"] = off
    # pictures -> card: anchor column inside the card, anchor row near the title (<=3 rows above .. last row);
    # each picture goes to the card whose title row is nearest
    owner = {}
    for k, im in enumerate(imgs):
        fr = getattr(im.anchor, "_from", None)
        if fr is None:
            continue
        col, row = fr.col + 1, fr.row + 1
        cand = []
        for card in cards:
            c0, c1, r0, r1 = card["_rect"]
            top, bot = card["_off"] + r0, card["_off"] + r1
            if c0 <= col <= c1 and top - 3 <= row <= bot:
                cand.append((abs(row - top), card["id"]))
        if cand:
            owner[k] = min(cand)[1]
    by_card = {}
    for k, cid in owner.items():
        by_card.setdefault(cid, []).append(k)

    def area(k):
        e = getattr(imgs[k].anchor, "ext", None)
        return (e.width * e.height) if e is not None else 0

    for card in cards:
        ks = sorted(by_card.get(card["id"], []), key=area, reverse=True)
        if not ks:
            continue
        found = {"photo": [], "drawing": []}
        loaded = []
        for k in ks:
            im = imgs[k]
            pil = None
            try:
                from PIL import Image
                pil = Image.open(io.BytesIO(im._data())).convert("RGBA")
                l, t, r, b = crop_of(im)
                if l or t or r or b:
                    w, h = pil.size
                    pil = pil.crop((round(l * w), round(t * h), round(w - r * w), round(h - b * h)))
            except Exception:
                pass
            loaded.append((k, pil))
        if len(loaded) == 1:  # a single picture is always the photo slot
            found["photo"].append(loaded[0])
        else:  # relative: the picture with the lowest tonal variety is the drawing (if clearly line-art-like)
            var = [(_variety(pil) if pil is not None else 99) for _, pil in loaded]
            lo = min(range(len(loaded)), key=lambda i: var[i])
            is_draw = loaded[lo][1] is not None and (_classify(loaded[lo][1]) == "drawing" or var[lo] <= 0.7 * max(var))
            for i, item in enumerate(loaded):
                found["drawing" if (i == lo and is_draw) else "photo"].append(item)
        for kind, key in (("photo", "image"), ("drawing", "drawing")):
            if not found[kind]:
                card.setdefault(key, None)
                continue
            if len(found[kind]) > 1:
                card.setdefault("_warnings", []).append(f"{len(found[kind])} {kind}s in card, kept the largest")
            k, pil = found[kind][0]
            ext = getattr(imgs[k].anchor, "ext", None)
            info = {"path": None, "status": "missing"}
            if img_dir and pil is not None:
                try:
                    from pl_images import prep_image
                    raw = img_dir / f"{card['id']}.{kind}.raw.png"
                    pil.save(raw)
                    out = img_dir / (f"{card['id']}.png" if kind == "photo" else f"{card['id']}.drawing.png")
                    prep_image(str(raw), str(out), remove_bg=False)
                    raw.unlink()
                    info = {"path": str(out), "status": "ok"}
                except Exception:
                    pass
            if ext is not None and ext.width:
                info["max_w_px"] = round(ext.width / 9525)
                info["max_h_px"] = round(ext.height / 9525)
            card[key] = info
    doc_prices, seen = [], set()
    for card in cards:
        card.pop("_off")
        if not keep_rect:
            card.pop("_rect")
        if reflow:
            card.pop("place", None)
            card["width"] = "auto"
            for c in card["columns"]:
                c.pop("span", None)
        okey = next((c["key"] for c in card["columns"] if column_role(c) == "order_code"), None)
        for row in card["rows"]:
            oc = row.get(okey) if okey else None
            if oc and oc not in seen:
                seen.add(oc)
                doc_prices.append({"order_code": oc, "price": prices.get(oc)})
    meta = {"brand": brand or "BRAND", "page_code": page_code or "BRD", "brand_logo": brand_logo}
    return {"meta": meta, "prices": doc_prices, "cards": cards}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--pages", default="1", help="comma list, 1-based, or 'all'")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--work", help="folder for extracted images (img/)")
    ap.add_argument("--brand-logo")
    ap.add_argument("--reflow", action="store_true", help="no locked places / fixed spans (width auto)")
    a = ap.parse_args()
    if a.pages == "all":
        ws = load_workbook(a.xlsx, read_only=True).worksheets[0]
        pages = list(range(1, -(-ws.max_row // load_theme()["page"]["rows"]) + 1))
    else:
        pages = [int(x) for x in a.pages.split(",")]
    doc = extract(a.xlsx, pages, a.work, a.brand_logo, reflow=a.reflow)
    save_doc(doc, a.out)
    print(f"cards={len(doc['cards'])} rows={sum(len(c['rows']) for c in doc['cards'])} "
          f"images={sum(1 for c in doc['cards'] if c['image'])} "
          f"drawings={sum(1 for c in doc['cards'] if c.get('drawing'))} "
          f"notes={sum(1 for c in doc['cards'] if c.get('note'))} prices={len(doc['prices'])}")


if __name__ == "__main__":
    main()
