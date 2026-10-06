"""Structural QA of a built price-list xlsx (optionally cross-checked against the placed doc)."""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

import pathlib as _pathlib, sys as _sys  # noqa: E401
_sys.path.insert(0, str(_pathlib.Path(__file__).resolve().parent))
from pl_common import col_index, col_letter, column_label, load_doc, load_theme


def _cf(s) -> str:
    return _norm(s).casefold()


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _scan_cards(ws, theme):
    """Find card headers by the 'Mã Hãng' label and read their columns and data rows from the sheet."""
    P, T = theme["page"], theme["table"]
    merges = {(r.min_row, r.min_col): r for r in ws.merged_cells.ranges}
    code_label = _cf(T["labels"]["code"])
    order_label = _cf(T["labels"]["order_code"])
    found = []
    for row in ws.iter_rows(min_col=P["body_left"], max_col=P["body_right"]):
        for cell in row:
            if _cf(cell.value) != code_label:
                continue
            rel = (cell.row - 1) % P["rows"] + 1
            if not (P["body_top"] <= rel <= P["body_bottom"]):
                continue
            labels, c = [], cell.column
            hh = merges[(cell.row, c)].max_row - cell.row + 1 if (cell.row, c) in merges else 1
            while c <= P["cols"]:
                v = ws.cell(cell.row, c).value
                if v is None:
                    break
                m = merges.get((cell.row, c))
                w = m.max_col - m.min_col + 1 if m else 1
                labels.append((_norm(v), c, w))
                c += w
            oc = next((c for lb, c, _ in labels if lb.casefold() == order_label), None)
            pc = labels[-1][1] if labels else None
            rows = []
            if oc is not None:
                r, w0, hc = cell.row + hh, labels[0][2], cell.column
                while r <= ws.max_row:
                    m = merges.get((r, hc))
                    if ws.cell(r, hc).value is None or (m and m.max_col - m.min_col + 1 != w0):
                        break  # blank row, or the full-width note under the table
                    rows.append({"order": ws.cell(r, oc).value, "order_cell": ws.cell(r, oc),
                                 "price_cell": ws.cell(r, pc)})
                    r += m.max_row - m.min_row + 1 if m else 1
            found.append({"row": cell.row, "col": cell.column, "labels": labels, "rows": rows})
    return found


def qa(xlsx: str, doc: dict | None, theme: dict) -> dict:
    P, F, H, T = theme["page"], theme["footer"], theme["header"], theme["table"]
    errors, warnings = [], []
    wb = load_workbook(xlsx)
    ws = wb[theme["sheet1_name"]] if theme["sheet1_name"] in wb.sheetnames else wb.worksheets[0]
    s2 = wb[theme["sheet2"]["name"]] if theme["sheet2"]["name"] in wb.sheetnames else None

    R = P["rows"]
    if ws.max_row % R:
        errors.append(f"sheet has {ws.max_row} rows, not a multiple of {R}")
    pages = -(-ws.max_row // R)

    # ---- frame
    title_cell = H["title"]["range"].split(":")[0]
    pc_cell = F["page_code"]["range"].split(":")[0]
    band_cell = F["band"]["range"].split(":")[0]
    prefix = None
    def addr(a, off):
        m = re.match(r"([A-Z]+)(\d+)", a)
        return f"{m.group(1)}{int(m.group(2)) + off}"

    expect_title = None
    if doc:
        expect_title = doc["meta"].get("title") or H["title"]["format"].format(brand=doc["meta"]["brand"])
    for pg in range(pages):
        off = pg * R
        t = ws[addr(title_cell, off)].value
        if not t or (expect_title and t != expect_title):
            errors.append(f"page {pg + 1}: title is {t!r}" + (f", expected {expect_title!r}" if expect_title else ""))
        code = ws[addr(pc_cell, off)].value
        if doc:
            want = P["page_code_format"].format(code=doc["meta"]["page_code"], n=pg + 1)
        else:
            m = re.match(r"^(.+) - (\d+)$", str(code or ""))
            prefix = prefix or (m.group(1) if m else None)
            want = P["page_code_format"].format(code=prefix, n=pg + 1) if prefix else None
        if code != want:
            errors.append(f"page {pg + 1}: page code is {code!r}, expected {want!r}")
        fill = ws[addr(band_cell, off)].fill
        if fill.fgColor.rgb != theme["colors"][F["band"]["fill"]]:
            errors.append(f"page {pg + 1}: footer red band missing")
    vl = H["varin_logo"]
    bl = H["brand_logo"]
    anchors = [(im.anchor._from.col, im.anchor._from.row) for im in ws._images if hasattr(im.anchor, "_from")]
    for pg in range(pages):
        if (col_index(vl["col"]) - 1, pg * R + vl["row"] - 1) not in anchors:
            errors.append(f"page {pg + 1}: Varin logo missing")
        if (col_index(bl["col"]) - 1, pg * R + bl["row"] - 1) not in anchors:
            msg = f"page {pg + 1}: brand logo missing"
            (errors if doc and doc["meta"].get("brand_logo") else warnings).append(msg)
    ids = sorted(b.id for b in ws.row_breaks.brk)
    if ids != [R * k for k in range(1, pages)]:
        errors.append(f"row breaks {ids} != expected {[R * k for k in range(1, pages)]}")
    if ws.page_setup.paperSize != P["paper_size"] or ws.sheet_view.view != P["view"]:
        errors.append("print setup / view differs from theme")

    # ---- sheet2 prices
    prices = {}
    if s2 is None:
        errors.append("Sheet2 missing")
    else:
        for r in s2.iter_rows(min_row=2, max_col=2, values_only=True):
            if r[0] is not None:
                prices[str(r[0])] = r[1]

    # ---- cards from sheet
    cards = _scan_cards(ws, theme)
    labels = T["labels"]
    seen, data_rows = {}, 0
    for c in cards:
        lb = [x[0] for x in c["labels"]]
        where = f"card at {col_letter(c['col'])}{c['row']}"
        if lb[0] != _norm(labels["code"]) or len(lb) < 3 or lb[-2] != _norm(labels["order_code"]) \
                or lb[-1] != _norm(labels["price"]):
            errors.append(f"{where}: header labels {lb} do not start/end as required")
        for r in c["rows"]:
            data_rows += 1
            o, pcell = r["order_cell"], r["price_cell"]
            v = str(o.value)
            if not (isinstance(o.value, str) and re.fullmatch(r"\d{7}", v)) or o.number_format != T["order_format"]:
                errors.append(f"{o.coordinate}: order code {o.value!r} is not 7-digit text")
            want = T["price_formula"].format(cell=o.coordinate)
            if pcell.value != want:
                errors.append(f"{pcell.coordinate}: price cell is not VLOOKUP formula ({pcell.value!r})")
            if v in seen:
                warnings.append(f"duplicate order code {v} at {o.coordinate} and {seen[v]}")
            seen.setdefault(v, o.coordinate)
            if s2 is not None:
                if v not in prices:
                    errors.append(f"order code {v} ({o.coordinate}) not in Sheet2")
                elif not isinstance(prices[v], (int, float)):
                    warnings.append(f"order code {v}: non-numeric price {prices[v]!r}")

    # ---- doc cross-checks
    if doc:
        rects = []
        n_drawings = 0
        expected_rows = 0
        for card in doc["cards"]:
            g = card.get("_geo")
            if not g:
                errors.append(f"card {card['id']}: no _geo")
                continue
            expected_rows += len(card["rows"])
            top, bot = g["row"], g["row"] + g["height"] - 1
            left, right = g["col"], g["col"] + g["width"] - 1
            if top < P["body_top"] or bot > P["body_bottom"] or left < P["body_left"] or right > P["body_right"]:
                errors.append(f"card {card['id']}: outside body area (rows {top}-{bot}, cols {left}-{right})")
            if len(g["spans"]) != len(card["columns"]):
                errors.append(f"card {card['id']}: {len(g['spans'])} spans for {len(card['columns'])} columns")
            if sum(g["spans"]) != g["width"]:
                errors.append(f"card {card['id']}: spans do not sum to width")
            rects.append((card["id"], g["page"], left, right, top, bot))
            img = g.get("image")
            if img and not (left <= img["col"] <= right and top <= img["row"] <= bot):
                errors.append(f"card {card['id']}: image anchor outside card")
            drw = g.get("drawing")
            if drw:
                n_drawings += 1
                if not (left <= drw["col"] <= right and top <= drw["row"] <= bot):
                    errors.append(f"card {card['id']}: drawing anchor outside card")
            for w in card.get("_warnings", []):
                warnings.append(f"card {card['id']}: {w}")
            if (card.get("image") or {}).get("status") == "review":
                warnings.append(f"card {card['id']}: image needs review")
            if (card.get("image") or {}).get("status") == "missing":
                warnings.append(f"card {card['id']}: image missing")
        for i in range(len(rects)):
            for j in range(i + 1, len(rects)):
                a, b = rects[i], rects[j]
                if a[1] == b[1] and a[2] <= b[3] and b[2] <= a[3] and a[4] <= b[5] and b[4] <= a[5]:
                    errors.append(f"cards overlap: {a[0]} and {b[0]}")
        if expected_rows != data_rows:
            errors.append(f"row count mismatch: doc has {expected_rows}, sheet has {data_rows}")
        if len(cards) != len(doc["cards"]):
            errors.append(f"card count mismatch: doc has {len(doc['cards'])}, sheet has {len(cards)}")

    stats = {"pages": pages, "cards": len(cards), "data_rows": data_rows, "images": len(ws._images),
             "sheet2_rows": len(prices)}
    if doc:
        stats["drawings"] = n_drawings
        u = uniformity(xlsx, doc, theme)
        errors += u["errors"]
        warnings += u["warnings"]
        stats["uniformity"] = {k: (sorted(v, key=str) if isinstance(v, set) else v) for k, v in u.items()
                               if k not in ("errors", "warnings")}
    return {"errors": errors, "warnings": warnings, "stats": stats}



def _med(v):
    v = sorted(v)
    n = len(v)
    return 0 if not n else (v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2)


def _rgb(color):
    return color.rgb if color is not None and isinstance(color.rgb, str) else None


def uniformity(xlsx: str, doc: dict, theme: dict) -> dict:
    """Cross-brand consistency metrics from the placed doc + the cells actually written."""
    P, C, T, U = theme["page"], theme["card"], theme["table"], theme["uniformity"]
    errors, warnings = [], []
    cards = [c for c in doc["cards"] if c.get("_geo")]
    first = [c for c in cards if not c["_geo"].get("cont")]  # continuation chunks carry no image by design
    body = (P["body_bottom"] - P["body_top"] + 1) * (P["body_right"] - P["body_left"] + 1)
    last = max((c["_geo"]["page"] for c in cards), default=0)
    area: dict[int, int] = {}
    for c in cards:
        g = c["_geo"]
        area[g["page"]] = area.get(g["page"], 0) + g["width"] * g["height"]
    fill = {p: round(a / body, 3) for p, a in sorted(area.items())}
    for p, f in fill.items():
        if p != last and f < U["min_fill"]:
            warnings.append(f"page {p + 1}: fill ratio {f} < {U['min_fill']}")

    heights = [c["_geo"]["image"]["cy_emu"] / 9525 for c in first if c["_geo"].get("image")]
    med = _med(heights)
    for c in first:
        im = c["_geo"].get("image")
        if not im:
            # user requirement 2026-10-06: every card must have a product image (theme image.required)
            (errors if theme.get("image", {}).get("required", True) else warnings).append(
                f"card {c['id']}: no image")
        elif im["cy_emu"] / 9525 < U["min_image_frac"] * med:
            warnings.append(f"card {c['id']}: image height {im['cy_emu'] / 9525:.0f}px < "
                            f"{U['min_image_frac']} x median {med:.0f}px")
    image_px = {"n": len(heights), "min": min(heights, default=0), "median": med, "max": max(heights, default=0)}
    desc: dict[int, int] = {}
    for c in cards:
        d = c["_geo"]["desc_rows"]
        desc[d] = desc.get(d, 0) + 1
    tl = [len(c["title"]) for c in cards]
    title_len = {"min": min(tl, default=0), "median": _med(tl), "max": max(tl, default=0)}
    labels: dict[str, int] = {}
    for c in cards:
        for col in c["columns"]:
            lb = _norm(column_label(col, theme))
            labels[lb] = labels.get(lb, 0) + 1

    colors = theme["colors"]
    ok_sizes = {float(x) for x in (C["title_size"], C["line_size"], C["vat_size"], C["note_size"],
                                   T["header_size"], T["data_size"])}
    ok = {"font": {theme["fonts"]["name"]}, "font size": ok_sizes, "font colour": {None, *colors.values()},
          "border colour": {colors["border"]}, "number format": {"General", T["order_format"], T["price_format"]}}
    used = {k: set() for k in ok}
    ws = load_workbook(xlsx).worksheets[0]
    for c in cards:
        g = c["_geo"]
        off = g["page"] * P["rows"]
        for r in range(g["row"] + off, g["row"] + off + g["height"]):
            for col in range(g["col"], g["col"] + g["width"]):
                cell = ws.cell(r, col)
                if cell.value is not None:
                    used["font"].add(cell.font.name)
                    used["font size"].add(float(cell.font.sz) if cell.font.sz else None)
                    used["font colour"].add(_rgb(cell.font.color))
                    used["number format"].add(cell.number_format)
                for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom):
                    if side is not None and side.style:
                        used["border colour"].add(_rgb(side.color))
    for what, vals in used.items():
        for x in sorted((v for v in vals if v not in ok[what]), key=str):
            errors.append(f"off-theme {what} used in card cells: {x!r}")
    return {"errors": errors, "warnings": warnings, "page_fill": fill, "image_px": image_px,
            "desc_rows": desc, "title_len": title_len, "labels": labels, "fonts": used["font"],
            "sizes": used["font size"], "colors": used["font colour"], "borders": used["border colour"],
            "formats": used["number format"]}


def compare_table(pairs, theme) -> tuple[str, bool]:
    """pairs = [(xlsx, placed_doc)]. Returns (text table, styles_equal)."""
    rows = []
    for x, d in pairs:
        u = uniformity(x, d, theme)
        fills = list(u["page_fill"].values())
        rows.append({"name": Path(x).name, "pages": len(fills), "cards": len(d["cards"]),
                     "rows": sum(len(c["rows"]) for c in d["cards"]),
                     "fill": sum(fills[:-1]) / len(fills[:-1]) if len(fills) > 1 else (fills[0] if fills else 0),
                     "last": fills[-1] if fills else 0,
                     "none": sum(1 for c in d["cards"] if not (c.get("_geo") or {}).get("image")
                                 and not (c.get("_geo") or {}).get("cont")), "img": u["image_px"]["median"],
                     "style": (frozenset(u["fonts"]), frozenset(u["sizes"]), frozenset(u["colors"]),
                               frozenset(u["borders"])),
                     "fmt": frozenset(u["formats"]), "errors": len(u["errors"])})
    s0, f0 = rows[0]["style"], rows[0]["fmt"]
    lines = [f"{'file':32} {'pages':>5} {'cards':>5} {'rows':>5} {'fill':>6} {'last_fill':>9} {'img_med':>7} {'img_none':>8} {'style=':>6} {'fmt=':>5} {'errs':>4}"]
    for r in rows:
        lines.append(f"{r['name'][:32]:32} {r['pages']:>5} {r['cards']:>5} {r['rows']:>5} {r['fill']:>6.2f} {r['last']:>9.2f} "
                     f"{r['img']:>7.0f} {r['none']:>8} {'Y' if r['style'] == s0 else 'N':>6} {'Y' if r['fmt'] == f0 else 'N':>5} {r['errors']:>4}")
    return "\n".join(lines), all(r["style"] == s0 and r["fmt"] == f0 for r in rows)


def write_report(result: dict, path) -> None:
    lines = ["# QA report", "", f"stats: {result['stats']}", "", f"## Errors ({len(result['errors'])})"]
    lines += [f"- {e}" for e in result["errors"]] or ["- none"]
    lines += ["", f"## Warnings ({len(result['warnings'])})"]
    lines += [f"- {w}" for w in result["warnings"]] or ["- none"]
    Path(path).write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="QA a built price-list xlsx")
    ap.add_argument("xlsx", nargs="?")
    ap.add_argument("--compare", nargs="+", metavar="XLSX:PLACED.JSON",
                    help="cross-brand table; exit 1 if style sets differ")
    ap.add_argument("--doc", help="placed doc.json to cross-check")
    ap.add_argument("-o", "--out", default="qa_report.md")
    a = ap.parse_args()
    if a.compare:
        pairs = []
        for item in a.compare:
            x, _, j = item.rpartition(":")
            pairs.append((x, load_doc(j)))
        text, ok = compare_table(pairs, load_theme())
        print(text)
        print("styles/formats equal across files" if ok else "STYLE SETS DIFFER")
        sys.exit(0 if ok else 1)
    if not a.xlsx:
        ap.error("xlsx or --compare required")
    doc = load_doc(a.doc) if a.doc else None
    theme = load_theme((doc or {}).get("meta", {}).get("theme_overrides"))
    res = qa(a.xlsx, doc, theme)
    write_report(res, a.out)
    print(f"errors={len(res['errors'])} warnings={len(res['warnings'])} stats={res['stats']} -> {a.out}")
    sys.exit(1 if res["errors"] else 0)


if __name__ == "__main__":
    main()
