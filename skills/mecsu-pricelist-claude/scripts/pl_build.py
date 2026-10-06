"""Draw a placed doc (cards with _geo) into an xlsx. Layout is not decided here."""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

import pathlib as _pathlib, sys as _sys  # noqa: E401
_sys.path.insert(0, str(_pathlib.Path(__file__).resolve().parent))
from pl_common import (column_label, column_role, col_index, col_letter, load_doc, load_theme,
                       norm_order_code, px_to_emu, resolve_path)


def _font(theme, size, bold=False, italic=False, color=None):
    return Font(name=theme["fonts"]["name"], size=size, bold=bold, italic=italic,
                color=theme["colors"][color] if color else None)


def _merge(ws, c0, r0, c1, r1):
    # overlapping cards are a layout bug reported by pl_qa; never crash the build over it
    new = CellRange(min_col=c0, min_row=r0, max_col=c1, max_row=r1)
    if (c1 > c0 or r1 > r0) and all(new.isdisjoint(m) for m in ws.merged_cells.ranges):
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r1, end_column=c1)


def _border(ws, theme, c0, r0, c1, r1, sides):
    """sides: 'all' = every cell gets all four sides; else subset of 'lrtb' applied on the rectangle edge."""
    side = Side(style=theme["table"]["border_style"], color=theme["colors"]["border"])
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(r, c)
            b = cell.border
            allb = sides == "all"
            cell.border = Border(
                left=side if allb or ("l" in sides and c == c0) else b.left,
                right=side if allb or ("r" in sides and c == c1) else b.right,
                top=side if allb or ("t" in sides and r == r0) else b.top,
                bottom=side if allb or ("b" in sides and r == r1) else b.bottom)


def _put(ws, c0, r0, c1, r1, value, font, h, v, wrap=False, fill=None, fmt=None):
    _merge(ws, c0, r0, c1, r1)
    cell = ws.cell(r0, c0)
    if isinstance(cell, MergedCell):
        return cell
    cell.value = value
    cell.font = font
    cell.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if fmt:
        cell.number_format = fmt
    if fill:
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                ws.cell(r, c).fill = fill
    return cell


def _anchor(ws, path, col0, row0, coff, roff, cx, cy, warnings, what):
    """col0/row0 are 0-based. Returns True when added."""
    try:
        img = Image(str(path))
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=col0, colOff=coff, row=row0, rowOff=roff),
                                   ext=XDRPositiveSize2D(cx, cy))
        ws.add_image(img)
        return True
    except Exception as e:  # missing / corrupt file
        warnings.append(f"{what}: image not drawn ({path}): {type(e).__name__}")
        return False


def _frame(ws, theme, doc, page, warnings, base):
    P = theme["page"]
    off = page * P["rows"]
    meta = doc["meta"]
    H, F = theme["header"], theme["footer"]

    def rng(s):
        c0, r0, c1, r1 = range_boundaries(s)
        return c0, r0 + off, c1, r1 + off

    t = H["title"]
    title = meta.get("title") or t["format"].format(brand=meta["brand"])
    _put(ws, *rng(t["range"]), title, _font(theme, t["size"], t["bold"]), t["align"], "center")
    v = H["vat_notice"]
    _put(ws, *rng(v["range"]), v["text"], _font(theme, v["size"], italic=v["italic"]), v["align"], "center")
    b = F["band"]
    _put(ws, *rng(b["range"]), b["text"], _font(theme, b["size"], b["bold"], color=b["color"]),
         b["align"], "center", fill=PatternFill("solid", fgColor=theme["colors"][b["fill"]]))
    for key in ("order_link", "order_hint"):
        f = F[key]
        _put(ws, *rng(f["range"]), f["text"], _font(theme, f["size"], italic=f["italic"], color=f["color"]),
             f["align"], "center")
    pc = F["page_code"]
    _put(ws, *rng(pc["range"]), P["page_code_format"].format(code=meta["page_code"], n=page + 1),
         _font(theme, pc["size"], pc["bold"], pc["italic"]), pc["align"], "center")

    vl = H["varin_logo"]
    _anchor(ws, resolve_path(vl["path"], base), col_index(vl["col"]) - 1, off + vl["row"] - 1,
            vl["col_off_emu"], vl["row_off_emu"], vl["cx_emu"], vl["cy_emu"], warnings, "varin logo")
    bl = H["brand_logo"]
    p = resolve_path(meta.get("brand_logo"), base)
    if p is None:
        return
    try:
        from PIL import Image as PI
        with PI.open(p) as im:
            w, h = im.size
        k = min(bl["max_cx_emu"] / px_to_emu(w), bl["max_cy_emu"] / px_to_emu(h))
        cx, cy = int(px_to_emu(w) * k), int(px_to_emu(h) * k)
    except Exception as e:
        warnings.append(f"brand logo not drawn ({p}): {type(e).__name__}")
        return
    _anchor(ws, p, col_index(bl["col"]) - 1, off + bl["row"] - 1, bl["col_off_emu"], bl["row_off_emu"],
            cx, cy, warnings, "brand logo")


def _card(ws, theme, card, warnings, base):
    g = card["_geo"]
    C, T = theme["card"], theme["table"]
    off = g["page"] * theme["page"]["rows"]
    c0, r0 = g["col"], g["row"] + off
    width, height, spans = g["width"], g["height"], g["spans"]
    desc = g["desc_rows"]
    img = g.get("image")
    note = card.get("note")
    cid = card["id"]

    text_w = width
    if g.get("title_cols"):
        text_w = g["title_cols"]
    elif img and img["row"] - g["row"] < C["title_rows"]:  # image beside the title
        text_w = max(2, min(width, img["col"] - g["col"]))
    _put(ws, c0, r0, c0 + text_w - 1, r0 + C["title_rows"] - 1, card["title"],
         _font(theme, C["title_size"], C["title_bold"]), C["title_halign"], C["title_valign"], wrap=True)

    texts = list(card.get("lines") or []) + list(card.get("legend") or [])
    if len(texts) > desc:
        warnings.append(f"card {cid}: {len(texts) - desc} description line(s) do not fit desc_rows")
    for i, t in enumerate(texts[:desc]):
        cell = ws.cell(r0 + C["title_rows"] + i, c0 + C["line_indent_cols"])
        cell.value = t
        cell.font = _font(theme, C["line_size"])
        cell.alignment = Alignment(horizontal=C["line_halign"], vertical=C["line_valign"])

    cols = card["columns"]
    starts, x = [], c0
    for s in spans:
        starts.append(x)
        x += s
    vat_r = r0 + C["title_rows"] + desc
    hdr_r = vat_r + 1
    data_r0 = hdr_r + T["header_rows"]
    drh = card.get("data_row_height", 1)
    rh = g.get("row_heights") or [1] * len(card["rows"])
    tops = [data_r0 + sum(rh[:k]) * drh for k in range(len(rh) + 1)]
    order_i = next((i for i, c in enumerate(cols) if column_role(c) == "order_code"), None)

    for i, col in enumerate(cols):
        role = column_role(col)
        cs, ce = starts[i], starts[i] + spans[i] - 1
        if role == "price":
            _put(ws, cs, vat_r, ce, vat_r, C["vat_text"], _font(theme, C["vat_size"]), C["vat_halign"], "center")
        _put(ws, cs, hdr_r, ce, hdr_r + T["header_rows"] - 1, column_label(col, theme),
             _font(theme, T["header_size"], T["header_bold"]), "center", T["valign"], wrap=True)
        _border(ws, theme, cs, hdr_r, ce, hdr_r + T["header_rows"] - 1, "all")
        align = col.get("align") or T["align"][role if role in T["align"] else "spec"]
        for ri, row in enumerate(card["rows"]):
            r, hh = tops[ri], rh[ri] * drh
            fmt = None
            if role == "price":
                if order_i is None:
                    warnings.append(f"card {cid}: price column without order_code column")
                    val = None
                else:
                    val = T["price_formula"].format(cell=f"{col_letter(starts[order_i])}{r}")
                font = _font(theme, T["data_size"], italic=T["price_italic"], color=T["price_color"])
                fmt = T["price_format"]
            elif role == "order_code":
                raw = row.get(col["key"])
                val = norm_order_code(raw) or raw
                font = _font(theme, T["data_size"], bold=T["order_bold"])
                fmt = T["order_format"]
            else:
                val = row.get(col["key"])
                font = _font(theme, T["data_size"])
            _put(ws, cs, r, ce, r + hh - 1, val, font, align, T["valign"], fmt=fmt, wrap=rh[ri] > 1)
            _border(ws, theme, cs, r, ce, r + hh - 1, "all")

    if note:
        nr0 = tops[-1]
        nr1 = nr0 + C["note_rows"] - 1
        _put(ws, c0, nr0, c0 + width - 1, nr1, note, _font(theme, C["note_size"], True, True), "left", "top", wrap=True)
        _border(ws, theme, c0, nr0, c0 + width - 1, nr1, "b")

    if C["outline"]:
        _border(ws, theme, c0, r0, c0 + width - 1, r0 + height - 1, "lrtb")

    drawn = [0, 0]
    for k, (geo_img, src) in enumerate(((img, card.get("image")), (g.get("drawing"), card.get("drawing")))):
        if not geo_img:
            continue
        what = "image" if k == 0 else "drawing"
        p = resolve_path((src or {}).get("path"), base)
        if p is None:
            warnings.append(f"card {cid}: {what} has no path")
        elif _anchor(ws, p, geo_img["col"] - 1, off + geo_img["row"] - 1, geo_img["col_off_emu"],
                     geo_img["row_off_emu"], geo_img["cx_emu"], geo_img["cy_emu"], warnings, f"card {cid} {what}"):
            drawn[k] = 1
    return tuple(drawn)


def _free(ws, theme, item, warnings, base):
    off = item["page"] * theme["page"]["rows"]
    c0, r0, c1, r1 = range_boundaries(item["range"])
    r0, r1 = r0 + off, r1 + off
    st = item.get("style") or {}
    if item["type"] == "text":
        _put(ws, c0, r0, c1, r1, item.get("text", ""),
             _font(theme, st.get("size", theme["card"]["line_size"]), st.get("bold", False), st.get("italic", False),
                   st.get("color")), st.get("align", "left"), st.get("valign", "center"), wrap=True)
    elif item["type"] == "image":
        p = resolve_path(item.get("path"), base)
        try:
            from PIL import Image as PI
            with PI.open(p) as im:
                w, h = im.size
            cx, cy = px_to_emu(st.get("w_px", w)), px_to_emu(st.get("h_px", h))
        except Exception as e:
            warnings.append(f"free image not drawn ({item.get('path')}): {type(e).__name__}")
            return
        _anchor(ws, p, c0 - 1, r0 - 1, 0, 0, cx, cy, warnings, "free image")
    else:
        warnings.append(f"free item type '{item['type']}' (shape) is not supported; ignored")


def build(doc: dict, out_path: str, theme: dict, base_dir=None) -> dict:
    base = Path(base_dir) if base_dir else None
    P = theme["page"]
    warnings: list[str] = []
    pages = max([doc["meta"].get("page_count") or 1] + [c["_geo"]["page"] + 1 for c in doc["cards"]]
                + [f["page"] + 1 for f in doc.get("free", [])])
    wb = Workbook()
    ws = wb.active
    ws.title = theme["sheet1_name"]
    for c in range(1, P["cols"] + 1):
        ws.column_dimensions[col_letter(c)].width = P["col_width"]
    for r in range(1, pages * P["rows"] + 1):
        ws.row_dimensions[r].height = P["row_height"]
    for pg in range(pages):
        _frame(ws, theme, doc, pg, warnings, base)
    images = drawings = 0
    for card in doc["cards"]:
        i, d = _card(ws, theme, card, warnings, base)
        images += i
        drawings += d
        warnings += [f"card {card['id']}: {w}" for w in card.get("_warnings", [])]
    for item in doc.get("free", []):
        _free(ws, theme, item, warnings, base)

    for k in range(1, pages):
        ws.row_breaks.append(Break(id=P["rows"] * k))
    ws.page_setup.paperSize = P["paper_size"]
    ws.page_setup.orientation = P["orientation"]
    ws.page_margins = PageMargins(**P["margins_in"])
    ws.sheet_view.view = P["view"]

    s2 = wb.create_sheet(theme["sheet2"]["name"])
    s2.append(theme["sheet2"]["headers"])
    for i, p in enumerate(doc.get("prices", []), start=2):
        code = norm_order_code(p["order_code"]) or p["order_code"]
        s2.cell(i, 1, code).number_format = theme["table"]["order_format"]
        s2.cell(i, 2, p.get("price"))
    wb.save(out_path)
    return {"pages": pages, "cards": len(doc["cards"]), "images": images, "drawings": drawings, "warnings": warnings}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="placed doc.json -> xlsx")
    ap.add_argument("doc")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()
    doc = load_doc(a.doc)
    theme = load_theme(doc["meta"].get("theme_overrides"))
    try:
        stats = build(doc, a.out, theme, base_dir=Path(a.doc).resolve().parent)
    except PermissionError:
        print(f"ERROR: cannot write {a.out} — file is open in Excel (or read-only). Close it and run again, "
              f"or pass another -o path.", file=sys.stderr)
        sys.exit(2)
    print(f"pages={stats['pages']} cards={stats['cards']} images={stats['images']} warnings={len(stats['warnings'])}")
    for w in stats["warnings"]:
        print("WARN", w)


if __name__ == "__main__":
    main()
