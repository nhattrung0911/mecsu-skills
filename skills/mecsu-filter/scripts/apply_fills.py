# -*- coding: utf-8 -*-
"""Tang 4 - buoc 2: gop moi nguon gia tri thanh file ket qua, kem tu kiem tra.

Nguon theo thu tu do tin: rule (const cum / suy tu ten) > corpus (phu thuoc ham)
> model (da doi chieu von tu + da kiem tra bang duong cong corpus).

Chot chan truoc khi ghi:
  - khong bao gio de len gia tri filter da co; chi dien o TRONG
  - gia tri model chi duoc nhan khi: dung dinh dang, confidence >= nguong, khong bi
    chinh ten san pham phu dinh, va khong tu mau thuan giua cac cau tra loi
  - corpus KHONG dong y thi chi danh dau REVIEW, khong loai (do duoc: 11/15 lan
    corpus bac oan gia tri dung)
  - key them vao phai la key nhom hang do that su dung
  - E cuoi cung khong duoc co key trung lap y het

Ghi xong tu mo lai file kiem tra; sai bat ky muc nao thi exit code 1.
"""
from __future__ import annotations

import argparse
import collections
import json
import pickle
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

from build_ai_queue import NameGuard

OUT_COLS = [
    ('description_chuan', 'C chuan hoa (de xuat)', 60),
    ('description_ly_do', 'Ly do sua ten', 28),
    ('canh_bao_ten', 'Mau thuan C <-> E', 46),
    ('filter_bo_sung', 'Filter bo sung (key: value)', 70),
    ('filter_nguon', 'Nguon tung gia tri', 60),
    ('filter_e_chuan', 'E chuan hoa (day du)', 90),
    ('filter_con_thieu', 'Con thieu - can nguoi/web', 40),
    ('filter_nghi_sai', 'Filter nghi thua/sai', 36),
    ('do_tin', 'Do tin', 10),
]


def parse_pairs(text):
    out = []
    for part in str(text).split('|'):
        part = part.strip()
        if ':' not in part:
            continue
        k, _, v = part.partition(':')
        k, v = k.strip(), v.strip()
        if k.upper() != 'N/A':
            out.append((k, v))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--input', type=Path, required=True)
    ap.add_argument('--audit', type=Path, required=True)
    ap.add_argument('--queue', type=Path)
    ap.add_argument('--corpus-filled', type=Path)
    ap.add_argument('--ai-values', type=Path, nargs='*', default=[])
    ap.add_argument('--ai-verify', type=Path)
    ap.add_argument('--out-dir', type=Path, required=True)
    ap.add_argument('--min-confidence', type=float, default=0.8)
    args = ap.parse_args()

    res, _ginfo = pickle.load(open(args.audit, 'rb'))
    guard = NameGuard(res)

    # --- gom gia tri tu cac nguon ---
    fills = collections.defaultdict(dict)      # row -> {attr: (value, source)}
    for row_index, row in enumerate(res):
        for attr, value in parse_pairs(row['filter_bo_sung']):
            fills[row_index][attr] = (value, 'rule')

    if args.corpus_filled and args.corpus_filled.exists():
        for item in json.loads(args.corpus_filled.read_text(encoding='utf-8')):
            source = 'corpus(%s, %s dong)' % ('+'.join(item['det']), item['rows'])
            fills[item['row']].setdefault(item['attr'], (item['value'], source))

    verdicts = {}
    if args.ai_verify and args.ai_verify.exists():
        for item in json.loads(args.ai_verify.read_text(encoding='utf-8')):
            verdicts[item.get('key') or item['ask_id']] = item['verify']

    ai_stats = collections.Counter()
    if args.queue and args.ai_values:
        asks = {a.get('key_hash', str(a['ask_id'])): a
                for a in json.loads(args.queue.read_text(encoding='utf-8'))['asks']}
        answers = {}
        for path in args.ai_values:
            if path.exists():
                for k, v in json.loads(path.read_text(encoding='utf-8')).items():
                    answers.setdefault(k, []).append(v)
        for ask_key, variants in answers.items():
            ask = asks.get(ask_key)
            if not ask:
                continue
            values = {str(v.get('value') or '').strip() for v in variants}
            if len(values) > 1:
                ai_stats['bat dong giua cac model'] += 1
                continue
            value = values.pop()
            confidence = min(float(v.get('confidence') or 0) for v in variants)
            if not value:
                ai_stats['model bo trong'] += 1
                continue
            if confidence < args.min_confidence:
                ai_stats['confidence thap'] += 1
                continue
            verdict = verdicts.get(ask_key)
            # `OUT_OF_RANGE` KHONG con tu dong loai. Do duoc tren 15 ca bi bac:
            # it nhat 11 ca la bac OAN gia tri DUNG. Duong cong doi chieu chi co mot
            # truc (`Size Ren`), trong khi `Chieu Dai Ren` con phu thuoc DAI CHIEU DAI
            # cua bulong; corpus khong co bulong M22 dai nen moi gia tri dai deu bi
            # coi la ngoai khoang. Do chinh xac 4/15 -> pha nhieu hon cuu.
            # Giu lai gia tri, danh dau REVIEW de nguoi duyet xem.
            if verdict == 'OUT_OF_RANGE':
                ai_stats['corpus khong dong y (van giu, danh dau REVIEW)'] += 1
            # Tu mau thuan: model tra hai gia tri khac nhau cho cung mot co, cung
            # tieu chuan. It nhat mot cai sai va khong biet cai nao -> bo ca hai.
            if verdict == 'AI_TU_MAU_THUAN':
                ai_stats['model tu mau thuan'] += 1
                continue
            # Von tu enum co the ep model chon mot gia tri KHONG dung: nhom U-Bolts
            # khong co "Nhung Nong Kem" nen model tra "Ma Kem" cho san pham ma ten
            # ghi ro la nhung nong. Chinh ten la bang chung manh hon von tu.
            contradicted = [i for i in ask['rows']
                            if guard.rejects(res[i], ask['attr'], value)]
            if contradicted:
                ai_stats['ten san pham phu dinh'] += 1
                continue
            source = 'ai(conf %.2f, %s)' % (confidence, verdict or 'NO_CHECK')
            for row_index in ask['rows']:
                fills[row_index].setdefault(ask['attr'], (value, source))
            ai_stats['nhan'] += 1

    # key hop le cua tung nhom, de chan key la
    valid_keys = collections.defaultdict(set)
    for row in res:
        if row['group']:
            for attr, _v in parse_pairs(row['filter_goc']):
                valid_keys[row['group']].add(attr)

    # --- ghi file ---
    wb_in = openpyxl.load_workbook(args.input, data_only=True)
    ws_in = wb_in[wb_in.sheetnames[0]]
    src_rows = list(ws_in.iter_rows(values_only=True))
    header = list(src_rows[0])
    body = src_rows[1:]
    if len(body) != len(res):
        raise SystemExit('audit %d dong nhung file goc %d dong' % (len(res), len(body)))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'chuan_hoa'
    ws.append(header + [label for _k, label, _w in OUT_COLS])
    ws.freeze_panes = 'A2'

    counts = collections.Counter()
    for row_index, (src, row) in enumerate(zip(body, res)):
        existing = parse_pairs(row['filter_goc'])
        have = {k for k, _v in existing}
        added, sources = [], []
        for attr, (value, source) in fills.get(row_index, {}).items():
            if attr in have:
                continue                                   # khong de len gia tri co san
            if row['group'] and valid_keys[row['group']] and attr not in valid_keys[row['group']]:
                counts['key la bi chan'] += 1
                continue
            added.append((attr, value))
            sources.append('%s <- %s' % (attr, source))
            have.add(attr)
        merged = existing + added
        still = [t.strip() for t in str(row['can_tra_cuu']).split(',')
                 if t.strip() and t.strip() not in dict(added)]

        if added:
            counts['dong duoc bo sung'] += 1
        counts['o duoc bo sung'] += len(added)
        flagged = any('OUT_OF_RANGE' in s or 'DISAGREE' in s for s in sources)
        confidence = ('REVIEW' if (still or row['mau_thuan'] or flagged)
                      else 'HIGH' if all('rule' in s or 'corpus' in s for s in sources) or not sources
                      else 'MEDIUM')
        ws.append(list(src) + [
            row['ten_de_xuat'], row['loi_ten'], row['mau_thuan'],
            ' | '.join('%s: %s' % kv for kv in added),
            ' | '.join(sources),
            ' | '.join('%s: %s' % kv for kv in merged) if added else '',
            ', '.join(still), row['filter_nghi_thua'], confidence,
        ])

    for i, width in enumerate([12, 18, 60, 8, 70] + [w for _k, _l, w in OUT_COLS], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.auto_filter.ref = 'A1:%s%d' % (
        openpyxl.utils.get_column_letter(len(header) + len(OUT_COLS)), len(body) + 1)

    args.out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    out_path = args.out_dir / ('%s_chuanhoa_%s.xlsx' % (args.input.stem, stamp))
    wb.save(out_path)

    # --- tu kiem tra bang code khac voi code vua ghi ---
    problems = []
    check = openpyxl.load_workbook(out_path, data_only=True, read_only=True)
    sheet = check[check.sheetnames[0]]
    seen_rows = 0
    n_src = len(header)
    for i, values in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            continue
        seen_rows += 1
        if list(values[:n_src]) != list(body[i - 1]):
            problems.append('dong %d: cot goc bi doi' % (i + 1))
        merged = values[n_src + 5]
        if merged:
            # Key lap voi value KHAC nhau la filter multi-value hop le
            # (`Nganh Hang: To Lap Rap | Nganh Hang: CNC`). Chi lap y het moi la loi.
            pairs = [p.strip() for p in str(merged).split('|')]
            dup = [p for p, c in collections.Counter(pairs).items() if c > 1]
            if dup:
                problems.append('dong %d: filter trung y het: %s' % (i + 1, dup[:3]))
    check.close()
    if seen_rows != len(body):
        problems.append('so dong ra %d != vao %d' % (seen_rows, len(body)))

    print('nguon AI: %s' % (dict(ai_stats) or '(khong dung AI)'))
    print('ket qua : %s' % dict(counts))
    print('-> %s' % out_path)
    if problems:
        print('\nTU KIEM TRA THAT BAI:')
        for p in problems[:20]:
            print('  ' + p)
        sys.exit(1)
    print('tu kiem tra: OK (cot goc nguyen ven, khong key trung, du dong)')


if __name__ == '__main__':
    main()
