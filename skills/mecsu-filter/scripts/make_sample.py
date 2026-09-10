# -*- coding: utf-8 -*-
"""Trich mot bo mau de do dac truoc khi chay ca file.

Lay TRON CUM (cac dong cung skeleton ten) chu khong lay ngau nhien tung dong:
schema filter duoc hoc tu peer cung cum, cat doi cum thi so do duoc khong con
phan anh dung lan chay that. Duyet vong tron qua cac nhom hang de bo mau khong
bi mot nhom lon chiem het.
"""
from __future__ import annotations

import argparse
import sys

if hasattr(sys.stdout, 'reconfigure'):      # ten san pham co dau, console Windows la cp1252
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
import collections
import re
from pathlib import Path

import openpyxl

from filter_audit import detect_columns

NUM = re.compile(r'\d+(?:[.,]\d+)?')
GRP = re.compile(r'^(.*?)\s+(?:Of|Của)\s+(.+)$', re.I)


def skeleton(name):
    return NUM.sub('#', re.sub(r'\s+', ' ', str(name).strip())).lower()


def group_of(cell):
    c = collections.Counter()
    for part in str(cell).split('|'):
        key = part.split(':')[0].strip()
        m = GRP.match(key)
        if m and m.group(2).strip().lower() != 'group home':
            c[m.group(2).strip()] += 1
    return c.most_common(1)[0][0] if c else '(khong ro)'


def columns(header, name_col, filter_col):
    """Vi tri cot ten + cot filter. Do tu header nhu moi script khac, khong dem cot cung.

    Hardcode `r[2]`/`r[4]` chay dung tren dung mot layout: tren Handtools, cot 4 la
    `category_lv1` nen moi dong ve nhom '(khong ro)' va bo mau hong am tham. Khong doan
    ra thi DUNG LAI de nguoi chi cot, khong doan bua.
    """
    try:
        found = detect_columns(header, {'name': name_col, 'filt': filter_col})
    except SystemExit:
        listing = ', '.join('%d=%s' % (i, h) for i, h in enumerate(header))
        raise SystemExit(
            'Khong doan duoc cot ten / cot filter tu header nay:\n  %s\n'
            'Chi ro bang --name-col <so> va --filter-col <so> (0-based, cung quy uoc '
            'voi filter_audit.py).' % listing
        )
    return found['name'], found['filt']


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--input', type=Path, required=True)
    ap.add_argument('--output', type=Path, required=True)
    ap.add_argument('--rows', type=int, default=10000)
    ap.add_argument('--name-col', type=int, help='ep cot ten (0-based)')
    ap.add_argument('--filter-col', type=int, help='ep cot filter (0-based)')
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.input, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(min_row=1, values_only=True)
    header = list(next(it))
    rows = [list(r) for r in it]
    wb.close()

    name_at, filt_at = columns(header, a.name_col, a.filter_col)

    clusters = collections.OrderedDict()
    for r in rows:
        clusters.setdefault(skeleton(r[name_at]), []).append(r)
    by_group = collections.OrderedDict()
    for sk, members in clusters.items():
        by_group.setdefault(group_of(members[0][filt_at]), []).append((sk, members))
    for g in by_group:                      # cum lon truoc trong moi nhom
        by_group[g].sort(key=lambda x: -len(x[1]))

    picked, taken, cursor = [], set(), {g: 0 for g in by_group}
    order = list(by_group)
    while len(picked) < a.rows:
        progressed = False
        for g in order:
            i = cursor[g]
            if i >= len(by_group[g]):
                continue
            sk, members = by_group[g][i]
            cursor[g] += 1
            progressed = True
            if sk in taken:
                continue
            taken.add(sk)
            picked.extend(members)
            if len(picked) >= a.rows:
                break
        if not progressed:
            break

    out = openpyxl.Workbook(write_only=True)
    sh = out.create_sheet(ws.title)
    sh.append(header)
    for r in picked:
        sh.append(r)
    a.output.parent.mkdir(parents=True, exist_ok=True)
    out.save(a.output)

    gc = collections.Counter(group_of(r[filt_at]) for r in picked)
    print(f'{len(picked)} dong / {len(taken)} cum / {len(gc)} nhom  -> {a.output}')
    for g, n in gc.most_common(10):
        print(f'  {n:6d}  {g}')


if __name__ == '__main__':
    main()
