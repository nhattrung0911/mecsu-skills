# -*- coding: utf-8 -*-
"""Vong 2 - rut thong so tu trang web, va dung file Excel giao cho nguoi dung.

Trong tam la `co_trong_nguon`: luat neo bang chung. Rieng cho do da sinh ba loi
lien tiep 2026-09-11, moi loi deu la BAC OAN mot gia tri dung.
"""
import importlib.util
import json
import subprocess
import sys
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'


def _nap(ten_file: str, ten_module: str):
    spec = importlib.util.spec_from_file_location(ten_module, SCRIPTS / ten_file)
    module = importlib.util.module_from_spec(spec)
    sys.modules[ten_module] = module
    spec.loader.exec_module(module)
    return module


_env = _nap('skill_env.py', 'vong2rut_env')
sys.modules['skill_env'] = _env
web = _nap('extract_from_web.py', 'vong2rut_web')
sheet = _nap('build_sheet.py', 'vong2rut_sheet')
del sys.modules['skill_env']


def _don(text: str) -> str:
    return web.don(text)


# --- luat neo bang chung ----------------------------------------------------

def test_con_so_bia_thi_bi_bac():
    """Chot chan cot loi: model tung bia so kem `confidence: 1.0`."""
    nguon = _don('Chiều dài 143 mm, đường kính 15 mm')
    assert not web.co_trong_nguon('999 mm', nguon)


def test_con_so_that_thi_giu_du_don_vi_viet_khac():
    """`143mm` va `143 mm` la mot - bac vi khac dau cach la bac oan."""
    nguon = _don('Chiều dài 143 mm')
    assert web.co_trong_nguon('143mm', nguon)


def test_chu_nhat_khong_bi_xoa_khi_chuan_hoa():
    """Loi that: loc bang [^a-z0-9] xoa sach `油性`, roi moi gia tri tieng Nhat
    deu bi bac oan du no nam ngay trong nguon."""
    assert _don('油性') != '', 'chu Nhat bi xoa khi chuan hoa'
    assert web.co_trong_nguon('油性', _don('タイプ 油性 マーカー'))


def test_ban_dich_kem_nguyen_van_thi_neo_duoc():
    """Loi that: model dich `油性` thanh `Gốc dầu (油性)`. So ca cum thi khong
    khop vi phan dich khong nam trong trang tieng Nhat - nhung nguyen van thi co."""
    nguon = _don('タイプ 油性')
    assert web.co_trong_nguon('Gốc dầu (油性)', nguon)


def test_ban_dich_khong_kem_nguyen_van_thi_bi_bac():
    """Script khong the xac minh mot ban dich tran. Bac la dung, va do la ly do
    prompt bat model phai kem nguyen van."""
    assert not web.co_trong_nguon('Gốc dầu', _don('タイプ 油性'))


def test_gia_tri_rong_thi_bi_bac():
    assert not web.co_trong_nguon('', _don('bat cu gi'))


def test_chu_bo_the_script_va_style():
    """Noi dung <script> lot vao text la rac gui cho model, tinh tien nhu chu that."""
    ra = web.chu('<style>a{color:red}</style><p>Dài 180mm</p><script>x=1</script>')
    assert 'Dài 180mm' in ra
    assert 'color' not in ra and 'x=1' not in ra


# --- dung file Excel giao ---------------------------------------------------

def _job(tmp_path: Path, codes, names=None, web_facts=None) -> Path:
    job = tmp_path / 'job'
    job.mkdir()
    (job / 'codes.json').write_text(json.dumps(codes, ensure_ascii=False), encoding='utf-8')
    if names is not None:
        (job / 'names.json').write_text(json.dumps(names, ensure_ascii=False), encoding='utf-8')
    if web_facts is not None:
        (job / 'web_facts.json').write_text(json.dumps(web_facts, ensure_ascii=False),
                                            encoding='utf-8')
    return job


def test_thong_so_san_co_trong_file_khong_bi_bo_quen(tmp_path):
    """Loi that: 11/20 san pham cua file that hien 'chua tim duoc thong so'
    trong khi thong so nam ngay trong file - vi chi lay tu vong 2."""
    job = _job(tmp_path,
               codes=[{'ma_goc': 'A6-4', 'dong': [3], 'dong_thong_so': ['Gồm 4 chi tiết']}],
               names=[{'ma': 'A6-4', 'ten_moi': 'Bộ đục 4 chi tiết', 'trang_thai': 'OK'}])
    gom = sheet.gom(job)
    assert gom['A6-4']['thong_so'], 'thong so san co trong file bi bo quen'
    assert 'Gồm 4 chi tiết' in str(gom['A6-4']['thong_so'])


def test_vong_2_bo_sung_thong_so_cho_dong_thieu(tmp_path):
    job = _job(tmp_path,
               codes=[{'ma_goc': 'PM-10BLA', 'dong': [4], 'dong_thong_so': []}],
               names=[{'ma': 'PM-10BLA', 'ten_moi': 'Bút sơn đen', 'trang_thai': 'OK'}],
               web_facts=[{'ma': 'PM-10BLA', 'url': 'https://x/y', 'ten_de_xuat': 'Bút sơn đen',
                           'thong_so': {'Màu sắc': 'Đen'}, 'trang_thai': 'OK', 'ghi_chu': ''}])
    gom = sheet.gom(job)
    assert gom['PM-10BLA']['thong_so'] == {'Màu sắc': 'Đen'}
    assert gom['PM-10BLA']['nguon'] == 'https://x/y'


def test_ten_cua_vong_1_thang_ten_do_web_de_xuat(tmp_path):
    """Vong 1 dua tren du lieu khach da co - do la can cu chac hon trang ban le."""
    job = _job(tmp_path,
               codes=[{'ma_goc': 'X', 'dong': [3], 'dong_thong_so': []}],
               names=[{'ma': 'X', 'ten_moi': 'TÊN THEO QUY ƯỚC', 'trang_thai': 'OK'}],
               web_facts=[{'ma': 'X', 'url': 'u', 'ten_de_xuat': 'tên từ trang bán lẻ',
                           'thong_so': {}, 'trang_thai': 'OK', 'ghi_chu': ''}])
    assert sheet.gom(job)['X']['ten'] == 'TÊN THEO QUY ƯỚC'


def test_khong_vong_nao_co_ket_qua_thi_thoat_khac_0(tmp_path):
    """Giao mot file rong con te hon khong giao."""
    job = _job(tmp_path, codes=[{'ma_goc': 'X', 'dong': [3]}])
    goc = tmp_path / 'goc.xlsx'
    import openpyxl
    openpyxl.Workbook().save(goc)
    xong = subprocess.run(
        [sys.executable, str(SCRIPTS / 'build_sheet.py'), '--input', str(goc),
         '--job', str(job), '--out', str(tmp_path / 'ra.xlsx')],
        capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=90)
    assert xong.returncode != 0


def test_khong_ghi_de_file_goc(tmp_path):
    """File goc la du lieu that cua khach - ghi de len no la khong the hoan tac."""
    import openpyxl
    goc = tmp_path / 'goc.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'], ws['B1'] = 'Description', 'Mã'
    ws['A2'], ws['B2'] = 'Bút sơn', 'PM-1'
    wb.save(goc)
    truoc = goc.read_bytes()

    job = _job(tmp_path,
               codes=[{'ma_goc': 'PM-1', 'dong': [2], 'dong_thong_so': ['Màu: Đen']}],
               names=[{'ma': 'PM-1', 'ten_moi': 'Bút sơn đen PM-1', 'trang_thai': 'OK'}])
    ra = tmp_path / 'ra.xlsx'
    xong = subprocess.run(
        [sys.executable, str(SCRIPTS / 'build_sheet.py'), '--input', str(goc),
         '--job', str(job), '--out', str(ra)],
        capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=90)
    assert xong.returncode == 0, xong.stdout + xong.stderr
    assert goc.read_bytes() == truoc, 'file goc bi ghi de'
    assert ra.exists()
