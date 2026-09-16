# -*- coding: utf-8 -*-
"""B1: doc file, do cot ma, tach tien to noi bo ra hang + ma hang. 0 token.

    python extract_codes.py --input <file.xlsx> --out codes.json
    python extract_codes.py --input <file.xlsx> --code-col 1      # neu do khong ra

Khong doan duoc cot thi DUNG va in header that kem ten co, khong doan bua.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):      # console Windows la cp1252
    sys.stdout.reconfigure(encoding='utf-8')

# Tien to noi bo -> hang. Do tren file mau 2026-09-11, buoc DO xac nhan 3/3 dung.
# CHI ghi vao day cai da xac minh: `BOS` KHONG duoc doan la Bosi (co the la Bosch),
# `TSU`, `ASA`, `KST`... chua ai kiem -> de hang rong, van giu ma.
HANG = {
    'SAT': 'SATA',
    'BSI': 'Bosi',
    'ANE': 'Anex',
}

# Header co the la ten cot ma. Tim theo thu tu nay, khop dau tien thang.
UNG_VIEN = ('part_number', 'partnumber', 'part no', 'ma hang', 'ma_hang',
            'ma san pham', 'sku', 'model', 'code', 'ma')

# Cot mo ta / ten san pham. Can de gom dong noi tiep o bo cuc kieu khoi.
UNG_VIEN_MO_TA = ('description', 'mo ta', 'ten san pham', 'ten hang', 'ten', 'name')

TIEN_TO = re.compile(r'^([A-Za-z]{2,4})-(.+)$')


def chuan_hoa(gia_tri: str) -> str:
    """Gom mot ma ve dang so sanh duoc: NFKC, bo dau cach thua, x nhan thanh 'x'."""
    ma = unicodedata.normalize('NFKC', str(gia_tri)).strip()
    ma = ma.replace('×', 'x').replace('✕', 'x').replace('✖', 'x')
    return re.sub(r'\s+', ' ', ma)


def tach(ma_goc: str) -> dict:
    """Tach tien to noi bo. Giu NGUYEN phan con lai, ke ca gach ngang va hau to chu.

    `SAT-96552K` -> hang SATA, ma hang `96552K`  (hau to K la that, khong duoc bo)
    `ANE-ABRS-2110` -> hang Anex, ma hang `ABRS-2110`  (gach ngang la that)
    `ANE-No.3510` -> hang Anex, ma hang `3510`  (rieng `No.` la cach ghi catalogue)
    """
    ma = chuan_hoa(ma_goc)
    khop = TIEN_TO.match(ma)
    if not khop:
        return {'ma_goc': ma, 'tien_to': '', 'hang': '', 'ma_hang': ma}

    tien_to, con_lai = khop.group(1).upper(), khop.group(2)

    # CHI cat khi tien to nam trong danh sach DA XAC MINH. Truoc day cat moi cum
    # 2-4 chu cai truoc gach ngang, nen `PM-10BLA` (PM = Paint Marker, mot phan
    # cua ma hang) bi cat thanh `10BLA` va ma sai di thang vao file giao. Kiem
    # "ten co chua ma hang khong" cung khong bat duoc, vi no so voi `10BLA` -
    # chinh la ket qua da hong.
    if tien_to not in HANG:
        return {'ma_goc': ma, 'tien_to': '', 'hang': '', 'ma_hang': ma}

    hang = HANG[tien_to]
    if hang == 'Anex':
        con_lai = re.sub(r'^No\.?\s*', '', con_lai, flags=re.IGNORECASE)
    return {'ma_goc': ma, 'tien_to': tien_to, 'hang': hang, 'ma_hang': con_lai}


def doc_bang(duong_dan: Path, sheet: str | None) -> list[list]:
    if duong_dan.suffix.lower() == '.csv':
        with duong_dan.open(encoding='utf-8-sig', newline='') as f:
            return [dong for dong in csv.reader(f)]
    import openpyxl
    wb = openpyxl.load_workbook(duong_dan, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    return [list(dong) for dong in ws.iter_rows(values_only=True)]


def bo_dau(text: str) -> str:
    """`Mã` -> `ma`, `Số lượng` -> `so luong`.

    Khong bo dau truoc khi cat tu thi `Mã` bi cat thanh ['m'] - ky tu `ã` khong
    nam trong [a-z0-9] nen bi vut - va cot `Mã` khong bao gio duoc nhan ra.
    """
    khong_dau = unicodedata.normalize('NFD', text)
    return ''.join(c for c in khong_dau if unicodedata.category(c) != 'Mn')


def tu_khoa(text: str) -> list[str]:
    """Cat header thanh tu: `part_number` -> ['part', 'number'], `Mã` -> ['ma']."""
    return [t for t in re.split(r'[^a-z0-9]+', bo_dau(chuan_hoa(text or '')).lower()) if t]


def tim_header(bang: list[list]) -> int:
    """Header khong nhat thiet o dong dau.

    File that `Tao ma Hatok.xlsx` co dong 0 rong, header o dong 1. Cho rang header
    luon o dong 0 thi doc ra mot bang toan rong.
    """
    for i, dong in enumerate(bang[:20]):
        if sum(1 for o in dong if o is not None and str(o).strip()) >= 2:
            return i
    return 0


def do_cot_theo(header: list, ung_vien: tuple) -> int | None:
    """Khop theo TU, khong khop theo chuoi con.

    Khop chuoi con tung lam `"ma" in "gamma"` thanh True, va cot `gamma` bi nhan
    nham la cot ma - dung kieu doan bua ma luat so 5 cua repo cam.
    """
    cot_tu = [tu_khoa(o) for o in header]
    for ten in ung_vien:
        can = tu_khoa(ten)
        for i, tu in enumerate(cot_tu):
            # day tu cua ung vien phai xuat hien lien tiep trong day tu cua header
            if any(tu[j:j + len(can)] == can for j in range(len(tu) - len(can) + 1)):
                return i
    return None


def do_cot(header: list) -> int | None:
    return do_cot_theo(header, UNG_VIEN)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input', type=Path, required=True, help='File .xlsx hoac .csv.')
    parser.add_argument('--code-col', type=int, help='Chi so cot ma, dem tu 0.')
    parser.add_argument('--sheet', help='Ten sheet. Mac dinh lay sheet dau.')
    parser.add_argument('--out', type=Path, help='Ghi JSON ra day. Khong co thi in ra man hinh.')
    args = parser.parse_args()

    if not args.input.exists():
        raise SystemExit('khong thay file: %s' % args.input)

    bang = doc_bang(args.input, args.sheet)
    if len(bang) < 2:
        raise SystemExit('file khong co dong du lieu nao')

    dong_header = tim_header(bang)
    header, dong_du_lieu = bang[dong_header], bang[dong_header + 1:]
    cot = args.code_col if args.code_col is not None else do_cot(header)
    if cot is None:
        print('KHONG do duoc cot ma. Header that cua file:', file=sys.stderr)
        for i, o in enumerate(header):
            print('  [%d] %s' % (i, o), file=sys.stderr)
        raise SystemExit('Chay lai kem --code-col <so>, dem tu 0.')
    if cot >= len(header):
        raise SystemExit('--code-col %d vuot qua so cot cua file (%d)' % (cot, len(header)))

    cot_mo_ta = do_cot_theo(header, UNG_VIEN_MO_TA)

    # Gop ma trung: cung mot ma tra mot lan, giu lai so dong de doi chieu nguoc.
    #
    # Bo cuc KIEU KHOI: mot san pham = mot dong co ma, roi N dong chi co cot mo ta
    # la cac dong thong so cua no (`Tao ma Hatok.xlsx`: 20 san pham trong 81 dong).
    # Bo cuc mot-dong-mot-san-pham van chay qua day, chi la khong dong nao noi tiep.
    gom: dict[str, dict] = {}
    dang_xet: dict | None = None
    for so_dong, dong in enumerate(dong_du_lieu, start=dong_header + 2):
        gia_tri = dong[cot] if cot < len(dong) else None
        mo_ta = (str(dong[cot_mo_ta]).strip()
                 if cot_mo_ta is not None and cot_mo_ta < len(dong) and dong[cot_mo_ta] is not None
                 else '')

        if gia_tri is not None and str(gia_tri).strip():
            muc = tach(gia_tri)
            khoa = muc['ma_goc']
            if khoa in gom:
                gom[khoa]['dong'].append(so_dong)
                dang_xet = gom[khoa]
            else:
                muc.update({'dong': [so_dong], 'ten_file': mo_ta, 'dong_thong_so': []})
                gom[khoa] = muc
                dang_xet = muc
        elif mo_ta and dang_xet is not None:
            # Mot o co the chua nhieu dong thong so, ngan bang ky tu xuong dong
            # (Excel goi la wrap text). Khong tach ra thi ca khoi dinh lam mot chuoi.
            dang_xet['dong_thong_so'].extend(
                d.strip() for d in mo_ta.splitlines() if d.strip())

    ket_qua = list(gom.values())
    theo_hang = Counter(m['hang'] or '(chua biet)' for m in ket_qua)

    print('cot ma: [%d] %s' % (cot, header[cot]))
    print('dong du lieu: %d -> ma duy nhat: %d' % (len(dong_du_lieu), len(ket_qua)))
    for hang, n in theo_hang.most_common():
        print('  %-14s %5d  %5.1f%%' % (hang, n, n / len(ket_qua) * 100))

    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(ket_qua, ensure_ascii=False, indent=2), encoding='utf-8')
        print('da ghi %s' % args.out)
    else:
        print(json.dumps(ket_qua[:5], ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
